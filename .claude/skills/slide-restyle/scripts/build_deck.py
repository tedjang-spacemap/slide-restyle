# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "PyYAML>=6", "fonttools>=4"]
# ///
"""Build the restyled deck from a plan.json (deterministic; the plan holds every judgment).

    uv run build_deck.py work/x/plan.json [--variant srmc] [--out outputs/x.pptx] [--render]

Reads plan.json (+ its content.json and DESIGN.md tokens), starts from `meta.base_template`, and adds one
slide per plan slide (more when a slide has to be split). Writes:
  --out            the deck (default outputs/<source stem>.pptx)
  <out>.report.md  conversion report: variant, estimated layouts, overflow actions, manual slides, fonts
  work/x/build/    build.json (per-slide fit results) and image crops; with --render also png/ via render.sh

Text sizes are fitted with the real font files (Pillow): shrink down to `overflow.shrink` limits, then
split by rows (io_steps). Every text comes from the source shape named by its `src`; run colors are
mapped through `deck.color_map`. Schema: reference/plan-schema.md.
"""
from __future__ import annotations

import argparse
import copy
import json
import math
import re
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

from fontTools.ttLib import TTFont
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from design_tokens import load as load_tokens, provenance  # noqa: E402

EMU_IN = 914400
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
WS = re.compile(r"\s+")
SKILL_ROOT = Path(__file__).resolve().parents[1]


def qa(tag):
    return f"{{{A}}}{tag}"


def emu(v_in):
    return Emu(int(round(v_in * EMU_IN)))


# ---------------------------------------------------------------- fonts & measuring

FONT_DIRS = [Path(".fonts"), Path.home() / "Library/Fonts", Path("/Library/Fonts"),
             Path("/System/Library/Fonts/Supplemental"), Path("/System/Library/Fonts")]
# (family, bold, italic) -> candidate files. Only faces DESIGN.md uses; unknown families fall back to fc-match.
FONT_FILES = {
    ("Arial", False, False): ["Arial.ttf"], ("Arial", True, False): ["Arial Bold.ttf"],
    ("Arial", False, True): ["Arial Italic.ttf"], ("Arial", True, True): ["Arial Bold Italic.ttf"],
    ("Roboto Light", False, False): ["Roboto-Light.ttf"], ("Roboto Light", False, True): ["Roboto-LightItalic.ttf"],
    ("Roboto Light", True, False): ["Roboto-Medium.ttf"], ("Roboto Light", True, True): ["Roboto-MediumItalic.ttf"],
    ("Pretendard", False, False): ["Pretendard-Regular.otf"], ("Pretendard", True, False): ["Pretendard-Bold.otf"],
    ("Pretendard Light", False, False): ["Pretendard-Light.otf"], ("Pretendard Light", True, False): ["Pretendard-SemiBold.otf"],
}
_font_cache: dict = {}
missing_fonts: set = set()


def font_file(family, bold, italic):
    for key in ((family, bold, italic), (family, bold, False), (family, False, False)):
        for name in FONT_FILES.get(key, []):
            for d in FONT_DIRS:
                if (d / name).exists():
                    return d / name
    try:
        style = ":bold" if bold else ""
        out = subprocess.run(["fc-match", "-f", "%{file}", f"{family}{style}"], capture_output=True, text=True, timeout=10)
        if out.stdout:
            missing_fonts.add(f"{family}{' Bold' if bold else ''} → {Path(out.stdout).name} (fc-match)")
            return Path(out.stdout)
    except FileNotFoundError:
        pass
    raise SystemExit(f"font not found: {family} bold={bold}")


def pil_font(family, bold, italic):
    key = (family, bold, italic)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(str(font_file(family, bold, italic)), 1000)
    return _font_cache[key]


_cmap_cache: dict = {}


def covers(family, bold, italic, ch):
    key = (family, bold, italic)
    if key not in _cmap_cache:
        _cmap_cache[key] = set(TTFont(str(font_file(family, bold, italic)), fontNumber=0, lazy=True).getBestCmap())
    return ord(ch) in _cmap_cache[key]


def is_ea(ch):
    o = ord(ch)
    return (0x1100 <= o <= 0x11FF or 0x3000 <= o <= 0x9FFF or 0xAC00 <= o <= 0xD7AF
            or 0xF900 <= o <= 0xFAFF or 0xFF00 <= o <= 0xFFEF)


@dataclass
class Style:
    latin: str
    ea: str
    size: float
    bold: bool = False
    italic: bool = False
    color: str = "FFFFFF"
    caps: bool = False
    name: str = ""

    def scaled(self, size):
        s = copy.copy(self)
        s.size = size
        return s


@dataclass
class Run:
    text: str
    style: Style


@dataclass
class Para:
    runs: list
    align: str = "l"
    line_spacing: float = 1.0
    space_before: float = 0.0      # pt
    bullet: str | None = None
    bullet_color: str | None = None
    marL: float = 0.0              # in
    indent: float = 0.0            # in


def run_pieces(run):
    """Split a run into (text, family, bold, italic) pieces by script (latin vs ea font)."""
    st = run.style
    out, cur, cur_ea = [], "", None
    text = run.text.upper() if st.caps else run.text
    for ch in text:
        e = is_ea(ch) or (not ch.isspace() and not covers(st.latin, st.bold, st.italic, ch))
        if cur and e != cur_ea:
            out.append((cur, st.ea if cur_ea else st.latin))
            cur = ""
        cur += ch
        cur_ea = e
    if cur:
        out.append((cur, st.ea if cur_ea else st.latin))
    return [(t, fam, st.bold, st.italic, st.size) for t, fam in out]


def text_width_pt(text, family, bold, italic, size):
    return pil_font(family, bold, italic).getlength(text) * size / 1000.0


def line_height_pt(family, bold, italic, size):
    a, d = pil_font(family, bold, italic).getmetrics()
    return (a + d) * size / 1000.0


def tokens_of(para):
    """Break opportunities: after spaces (word wrap). Each token = list of (text, family, b, i, size)."""
    toks, cur = [], []
    for run in para.runs:
        for t, fam, b, i, sz in run_pieces(run):
            for part in re.split(r"(?<= )", t):
                if not part:
                    continue
                cur.append((part, fam, b, i, sz))
                if part.endswith(" "):
                    toks.append(cur)
                    cur = []
    if cur:
        toks.append(cur)
    return toks


def tok_width(tok, strip_trailing=False):
    w = 0.0
    for k, (t, fam, b, i, sz) in enumerate(tok):
        if strip_trailing and k == len(tok) - 1:
            t = t.rstrip(" ")
        w += text_width_pt(t, fam, b, i, sz)
    return w


def wrap_count(para, width_pt):
    """Number of lines the paragraph takes at width_pt (word wrap, long words broken by character)."""
    avail = width_pt - max(0.0, (para.marL + para.indent) * 72)
    if para.bullet:
        avail = width_pt - para.marL * 72
    lines, x = 1, 0.0
    for tok in tokens_of(para):
        w_full, w_trim = tok_width(tok), tok_width(tok, strip_trailing=True)
        if x + w_trim <= avail + 0.01:
            x += w_full
            continue
        if x > 0:
            lines += 1
            x = 0.0
        if w_trim <= avail:
            x = w_full
            continue
        for t, fam, b, i, sz in tok:          # word longer than the line: break by character
            for ch in t:
                cw = text_width_pt(ch, fam, b, i, sz)
                if x + cw > avail + 0.01 and x > 0:
                    lines += 1
                    x = 0.0
                x += cw
    return lines


def para_height(para, width_pt):
    lh = 0.0
    for run in para.runs or [Run(" ", Style("Arial", "Pretendard", 12))]:
        for t, fam, b, i, sz in run_pieces(run) or [(" ", run.style.latin, run.style.bold, run.style.italic, run.style.size)]:
            lh = max(lh, line_height_pt(fam, b, i, sz))
    return wrap_count(para, width_pt) * lh * para.line_spacing + para.space_before


def text_height(paras, width_pt):
    return sum(para_height(p, width_pt) for p in paras)


def max_lines(paras, width_pt):
    return max((wrap_count(p, width_pt) for p in paras), default=0)


# ---------------------------------------------------------------- design tokens

class Design:
    def __init__(self, tokens, variant):
        self.t = tokens
        self.variant = variant
        self.chrome = tokens["chrome"][variant]
        ov = self.chrome.get("style_overrides") or {}
        self.font_over = ov.get("fonts") or {}
        self.color_over = ov.get("colors") or {}
        self.colors = tokens["colors"]
        self.shrink = tokens["overflow"]["shrink"]

    def hex(self, name):
        if name is None:
            return None
        if re.fullmatch(r"[0-9A-Fa-f]{6}", str(name)):
            return str(name).upper()
        name = self.color_over.get(name, name)
        c = self.colors[name]
        return str(c.get("hex") or c["hexes"][0]).upper()

    def color(self, name):
        return self.colors[self.color_over.get(name, name)]

    def style(self, name, color=None, **kw):
        t = self.t["type"][name]
        role = self.font_over.get(t["font"], t["font"])
        f = self.t["fonts"][role]
        bold = t.get("bold", f.get("bold", False))
        col = color if color is not None else t.get("color", "text")
        st = Style(latin=f["latin"], ea=f["ea"], size=float(t["size"]), bold=bool(bold),
                   italic=bool(t.get("italic", False)), color=self.hex(col) if not str(col).startswith(("status.", "tag.")) else "FFFFFF",
                   caps=bool(t.get("caps", False)), name=name)
        for k, v in kw.items():
            setattr(st, k, v)
        return st

    def min_size(self, st):
        floor = (self.shrink.get("floors") or {}).get(st.name)
        lo = st.size * float(self.shrink.get("min_ratio", 0.85))
        if floor is not None:
            lo = max(lo, min(float(floor), st.size))
        return lo


class Geo:
    """Maps layout coordinates from the layout's native variant to the chosen one (variant_mapping)."""

    def __init__(self, design: Design, layout: dict):
        native = layout.get("native", "none")
        self.on = native not in ("none", design.variant)
        if self.on:
            self.N = design.t["chrome"][native]["content_area"]
            self.A = design.chrome["content_area"]

    def x(self, v):
        return self.A[0] + (v - self.N[0]) * self.A[2] / self.N[2] if self.on else v

    def y(self, v):
        return self.A[1] + (v - self.N[1]) * self.A[3] / self.N[3] if self.on else v

    def w(self, v):
        return v * self.A[2] / self.N[2] if self.on else v

    def h(self, v):
        return v * self.A[3] / self.N[3] if self.on else v

    def box(self, b):
        return [self.x(b[0]), self.y(b[1]), self.w(b[2]), self.h(b[3])]


# ---------------------------------------------------------------- source text

def norm(s):
    return WS.sub(" ", s or "").strip()


class Source:
    def __init__(self, content, plan, design: Design):
        self.shapes = {s["id"]: s for sl in content["slides"] for s in sl["shapes"]}
        self.slides = content["slides"]
        self.design = design
        deck = plan["deck"]
        self.tags = deck.get("tags", {})
        self.cmap = {k.upper(): v for k, v in (deck.get("color_map") or {}).items()}

    def tag_hex(self, tag):
        info = self.tags.get(tag)
        return self.design.hex(info["color"]) if info else self.design.hex("text_label")

    def mapped(self, hexv):
        v = self.cmap.get((hexv or "").upper())
        if v is None:
            return None
        if v.startswith("tag:"):
            return self.tag_hex(v[4:])
        return self.design.hex(v)

    def runs(self, ref):
        """Source runs covering ref["t"]: list of paragraphs, each a list of (text, src_run)."""
        sh = self.shapes[ref["src"]]
        paras = sh["text"]["paragraphs"]
        chars = []                                   # (char, run) over the whole shape, "\n" between paragraphs
        for pi, p in enumerate(paras):
            if pi:
                chars.append(("\n", None))
            for r in p["runs"]:
                chars.extend((ch, r) for ch in r["text"])
        plain = "".join(c for c, _ in chars)
        t = ref["t"]
        i = plain.find(t)
        n = len(t)
        if i < 0:
            m = re.search(r"\s+".join(re.escape(w) for w in t.split()), plain)
            if not m:
                return [[(t, None)]]
            i, n = m.start(), m.end() - m.start()
        out, cur, cur_run, buf = [], [], None, ""
        for ch, r in chars[i:i + n]:
            if ch == "\n":
                if buf:
                    cur.append((buf, cur_run))
                out.append(cur)
                cur, buf, cur_run = [], "", None
                continue
            if r is not cur_run and buf:
                cur.append((buf, cur_run))
                buf = ""
            cur_run = r
            buf += ch
        if buf:
            cur.append((buf, cur_run))
        out.append(cur)
        return out

    def paras(self, ref, style: Style, keep_runs=False, **pkw):
        """Paragraphs for a ref in `style`. keep_runs: carry the source run colour (via color_map) and bold."""
        out = []
        for pr in self.runs(ref):
            runs = []
            for text, r in pr:
                st = style
                if keep_runs and r is not None:
                    st = copy.copy(style)
                    hexv = self.mapped((r.get("color") or {}).get("hex"))
                    if hexv and hexv != self.design.hex("text"):
                        st.color = hexv
                    if r.get("bold") and not style.bold:
                        st.bold = True
                runs.append(Run(text, st))
            out.append(Para(runs, **pkw))
        return out

    def text(self, ref):
        return "\n".join("".join(t for t, _ in p) for p in self.runs(ref))


# ---------------------------------------------------------------- pptx writing

def strip_style(shape):
    st = shape._element.find(f"{{{P}}}style")
    if st is not None:
        shape._element.remove(st)


def set_fill(shape, hexv=None, alpha=None):
    if hexv is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hexv)
    if alpha is not None and alpha < 1:
        clr = shape.fill._xPr.find(f"{qa('solidFill')}/{qa('srgbClr')}")
        etree.SubElement(clr, qa("alpha")).set("val", str(int(round(alpha * 100000))))


def set_line(shape, hexv=None, width_pt=0.75, dash=None, alpha=None):
    ln = shape.line
    if hexv is None:
        ln.fill.background()
        return
    ln.color.rgb = _rgb(hexv)
    ln.width = Pt(width_pt)
    if alpha is not None and alpha < 1:
        clr = ln._get_or_add_ln().find(f"{qa('solidFill')}/{qa('srgbClr')}")
        etree.SubElement(clr, qa("alpha")).set("val", str(int(round(alpha * 100000))))
    if dash:
        lnel = ln._get_or_add_ln()
        pd = lnel.find(qa("prstDash"))
        if pd is None:
            pd = etree.SubElement(lnel, qa("prstDash"))
        pd.set("val", "dash" if dash is True else dash)


def _rgb(hexv):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexv.upper())


def rpr(parent, st: Style, lang="ko-KR"):
    r = etree.SubElement(parent, qa("rPr"))
    r.set("lang", lang)
    r.set("sz", str(int(round(st.size * 100))))
    r.set("b", "1" if st.bold else "0")
    r.set("i", "1" if st.italic else "0")
    if st.caps:
        r.set("cap", "all")
    r.set("dirty", "0")
    sf = etree.SubElement(r, qa("solidFill"))
    etree.SubElement(sf, qa("srgbClr")).set("val", st.color.upper())
    etree.SubElement(r, qa("latin")).set("typeface", st.latin)
    etree.SubElement(r, qa("ea")).set("typeface", st.ea)
    etree.SubElement(r, qa("cs")).set("typeface", st.latin)
    return r


def write_paras(tf, paras):
    body = tf._txBody
    for p in body.findall(qa("p")):
        body.remove(p)
    for para in paras:
        p = etree.SubElement(body, qa("p"))
        ppr = etree.SubElement(p, qa("pPr"))
        ppr.set("algn", {"l": "l", "ctr": "ctr", "r": "r", "just": "just"}.get(para.align, "l"))
        if para.marL or para.indent:
            ppr.set("marL", str(int(para.marL * EMU_IN)))
            ppr.set("indent", str(int(para.indent * EMU_IN)))
        if para.line_spacing != 1.0:
            ls = etree.SubElement(ppr, qa("lnSpc"))
            etree.SubElement(ls, qa("spcPct")).set("val", str(int(para.line_spacing * 100000)))
        sb = etree.SubElement(ppr, qa("spcBef"))
        etree.SubElement(sb, qa("spcPts")).set("val", str(int(round(para.space_before * 100))))
        if para.bullet:
            if para.bullet_color:
                bc = etree.SubElement(ppr, qa("buClr"))
                etree.SubElement(bc, qa("srgbClr")).set("val", para.bullet_color)
            font = para.runs[0].style.latin if para.runs else "Arial"
            etree.SubElement(ppr, qa("buFont")).set("typeface", font)   # else LibreOffice falls back to OpenSymbol
            etree.SubElement(ppr, qa("buChar")).set("char", para.bullet)
        else:
            etree.SubElement(ppr, qa("buNone"))
        for run in para.runs:
            for text, st in coverage_runs(run):
                r = etree.SubElement(p, qa("r"))
                rpr(r, st)
                etree.SubElement(r, qa("t")).text = text
        if para.runs:
            end = etree.SubElement(p, qa("endParaRPr"))
            end.set("lang", "ko-KR")
            end.set("sz", str(int(round(para.runs[-1].style.size * 100))))
            end.set("dirty", "0")


def coverage_runs(run):
    """Split off characters the latin font lacks (①, →) into runs whose latin face is the ea font —
    renderers pick the latin slot for them and would otherwise fall back to an arbitrary font."""
    st, out, cur, cur_alt = run.style, [], "", False
    for ch in run.text:
        alt = not ch.isspace() and not is_ea(ch) and not covers(st.latin, st.bold, st.italic, ch)
        if cur and alt != cur_alt:
            out.append((cur, cur_alt))
            cur = ""
        cur += ch
        cur_alt = alt
    if cur:
        out.append((cur, cur_alt))
    res = []
    for text, alt in out:
        s2 = st
        if alt:
            s2 = copy.copy(st)
            s2.latin = st.ea
        res.append((text, s2))
    return res


def body_pr(tf, insets=(0, 0, 0, 0), anchor="t", autofit="none", wrap=True):
    bp = tf._txBody.find(qa("bodyPr"))
    for ch in list(bp):
        bp.remove(ch)
    for k in list(bp.attrib):
        del bp.attrib[k]
    bp.set("wrap", "square" if wrap else "none")
    for k, v in zip(("lIns", "tIns", "rIns", "bIns"), insets):
        bp.set(k, str(int(round(v * EMU_IN))))
    bp.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}.get(anchor, "t"))
    bp.set("rtlCol", "0")
    etree.SubElement(bp, qa({"shape": "spAutoFit", "norm": "normAutofit"}.get(autofit, "noAutofit")))


# ---------------------------------------------------------------- builder

@dataclass
class Page:
    slide: object
    spec: dict
    out_n: int
    part: tuple | None
    fit: dict = field(default_factory=dict)


class Builder:
    def __init__(self, design: Design, plan, content, src: Source, base_dir: Path, build_dir: Path):
        self.d = design
        self.plan = plan
        self.content = content
        self.src = src
        self.base_dir = base_dir
        self.build_dir = build_dir
        self.root = SKILL_ROOT.parents[2] if (SKILL_ROOT.parents[2] / "design").exists() else Path.cwd()
        base = Path(design.t["meta"]["base_template"])
        self.prs = Presentation(str(base if base.exists() else self.root / base))
        self.theme_changes = replace_theme_fonts(self.prs, design)
        self.layout = next(l for m in self.prs.slide_masters for l in m.slide_layouts
                           if l.name == design.t["meta"]["base_layout"])
        self.pages: list[Page] = []
        self.box_used: dict = {}

    def asset(self, rel):
        p = Path(rel)
        return p if p.exists() else self.root / rel

    # -- primitives
    def rect(self, slide, box, fill=None, alpha=None, line=None, width=0.75, dash=None, shape=MSO_SHAPE.RECTANGLE, adj=None, line_alpha=None):
        sh = slide.shapes.add_shape(shape, emu(box[0]), emu(box[1]), emu(box[2]), emu(box[3]))
        strip_style(sh)
        if adj is not None:
            sh.adjustments[0] = adj / 100000
        set_fill(sh, fill, alpha)
        set_line(sh, line, width, dash, line_alpha)
        return sh

    def line(self, slide, x1, y1, x2, y2, hexv, width_pt, head=None):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        strip_style(c)
        c.line.color.rgb = _rgb(hexv)
        c.line.width = Pt(width_pt)
        if head:
            ln = c.line._get_or_add_ln()
            etree.SubElement(ln, qa("tailEnd")).set("type", head)
        return c

    def picture(self, slide, path, box, fit=None, align="center"):
        if fit == "contain":
            with Image.open(path) as im:
                iw, ih = im.size
            bw, bh = box[2], box[3]
            s = min(bw / iw, bh / ih)
            w, h = iw * s, ih * s
            x = box[0] + {"center": (bw - w) / 2, "top-left": 0, "top-right": bw - w}.get(align, (bw - w) / 2)
            y = box[1] + ((bh - h) / 2 if align == "center" else 0)
            box = [x, y, w, h]
        return slide.shapes.add_picture(str(path), emu(box[0]), emu(box[1]), emu(box[2]), emu(box[3]))

    def text(self, slide, box, paras, insets=(0, 0, 0, 0), anchor="t", autofit="none", name=None):
        tb = slide.shapes.add_textbox(emu(box[0]), emu(box[1]), emu(box[2]), emu(box[3]))
        if name:
            tb.name = name
        body_pr(tb.text_frame, insets, anchor, autofit)
        write_paras(tb.text_frame, paras)
        return tb

    def text_in(self, sh, paras, insets, anchor):
        body_pr(sh.text_frame, insets, anchor, "none")
        write_paras(sh.text_frame, paras)

    # -- slides
    def new_slide(self, spec, part=None, chrome=True):
        slide = self.prs.slides.add_slide(self.layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        page = Page(slide, spec, len(self.pages) + 1, part)
        self.pages.append(page)
        if chrome:
            self.draw_chrome(page)
        notes = []
        src_notes = self.content["slides"][spec["src"] - 1].get("notes")
        if src_notes:
            notes.append(src_notes)
        notes.append(f"원본 S{spec['src']} · 레이아웃 {spec['layout']}{'.' + spec['layout_variant'] if spec.get('layout_variant') else ''}")
        slide.notes_slide.notes_text_frame.text = "\n".join(notes)
        return page

    def kicker_text(self, spec, part):
        k = spec.get("slots", {}).get("kicker") or self.plan["deck"]["title"]
        t = self.src.text(k)
        sp = part or ((spec["split"]["part"], spec["split"]["of"]) if spec.get("split") else None)
        if sp:
            t += f" · {sp[0]}/{sp[1]}"
        return t

    def draw_chrome(self, page: Page):
        d, ch, slide, spec = self.d, self.d.chrome, page.slide, page.spec
        slots = spec.get("slots", {})
        defaults = ch.get("text_box_defaults", {})
        for bg in ch.get("background", []):
            if bg["kind"] == "rect":
                self.rect(slide, bg["box"], fill=d.hex(bg["fill"]))
            elif bg["kind"] == "picture":
                self.picture(slide, self.asset(bg["asset"]), bg["box"])
        has_status = bool(slots.get("status"))
        for el in ch["elements"]:
            kind, box = el["kind"], list(el["box"])
            if kind == "picture":
                self.picture(slide, self.asset(el["asset"]), box)
            elif kind == "line":
                self.line(slide, box[0], box[1], box[0] + box[2], box[1] + box[3], d.hex(el["color"]), el["width_pt"])
            elif kind == "pill":
                if has_status:
                    self.pill(slide, box, slots["status"], "header")
            elif kind == "text":
                slot = el["slot"]
                st = d.style(el["style"])
                insets = el.get("insets", defaults.get("insets", [0, 0, 0, 0]))
                anchor = defaults.get("anchor", "t")
                align = el.get("align", d.t["type"][el["style"]].get("align", "l"))
                if slot == "kicker":
                    txt = self.kicker_text(spec, page.part)
                    fmt = el.get("format", "{text}")
                    txt = fmt.replace("{TEXT}", txt).replace("{text}", txt)
                    if has_status and el.get("kicker_w_when_present"):
                        box[2] = el["kicker_w_when_present"]
                    self.fit_single_line(slide, box, txt, st, insets, anchor, align, page, "kicker")
                elif slot == "title":
                    if has_status:
                        box[2] = next((e.get("title_w_when_present") for e in ch["elements"] if e["id"] == "status" and e.get("title_w_when_present")), box[2])
                    self.draw_title(slide, box, spec, st, insets, anchor, page)
                elif slot == "footer_left":
                    self.text(slide, box, [Para([Run(self.plan["deck"].get("footer_left") or el.get("default", ""), st)], align=align)], insets, anchor)
                elif slot == "footer_right":
                    self.text(slide, box, [Para([Run(el.get("default", ""), st)], align=align)], insets, anchor)
                elif slot == "page_number":
                    self.text(slide, box, [Para([Run(f"{page.out_n:02d}", st)], align=align)], insets, anchor)
                elif slot == "source_line":
                    sl = slots.get("source_line")
                    if sl:
                        self.text(slide, box, self.src.paras(sl, st, align=align), insets, anchor)

    def fit_single_line(self, slide, box, txt, st, insets, anchor, align, page, what):
        width = (box[2] - insets[0] - insets[2]) * 72
        size, lo = st.size, self.d.min_size(st)
        while size > lo and max_lines([Para([Run(txt, st.scaled(size))])], width) > 1:
            size = max(lo, size - 0.5)
        para = Para([Run(txt, st.scaled(size))], align=align)
        if max_lines([para], width) > 1:
            page.fit.setdefault("issues", []).append(f"{what} 1줄 초과")
        if size < st.size:
            page.fit.setdefault("shrunk", {})[what] = round(size, 1)
        self.text(slide, box, [para], insets, anchor)

    def draw_title(self, slide, box, spec, st, insets, anchor, page):
        ref = spec.get("slots", {}).get("title")
        if not ref:
            return
        width = (box[2] - insets[0] - insets[2]) * 72
        paras = self.src.paras(ref, st)
        floor = float((self.d.shrink.get("floors") or {}).get("title", st.size))
        size = st.size
        while size > floor and max_lines(self.restyle(paras, size / st.size), width) > 1:
            size = max(floor, size - 1)
        out = self.restyle(paras, size / st.size)
        lines = max_lines(out, width)
        if lines > 1:          # reflow: two lines at the floor size, box grows downward (overflow.reflow.title)
            if self.d.variant == "collision" and lines == 2:
                box = [box[0], box[1], box[2], box[3] + 0.61]
                page.fit.setdefault("reflow", []).append("title 2줄")
            else:
                page.fit.setdefault("issues", []).append(f"제목 {lines}줄")
        if size < st.size:
            page.fit.setdefault("shrunk", {})["title"] = size
        self.text(slide, box, out, insets, anchor)

    @staticmethod
    def restyle(paras, ratio, mins=None):
        out = []
        for p in paras:
            q = copy.copy(p)
            q.runs = []
            for r in p.runs:
                lo = (mins or {}).get(r.style.name, 0)
                q.runs.append(Run(r.text, r.style.scaled(max(lo, round(r.style.size * ratio * 2) / 2))))
            out.append(q)
        return out

    def pill(self, slide, box, ref, where):
        d = self.d
        comp = d.t["components"]["pill"]
        status = d.color(ref.get("color", "status_blue"))
        st = d.style("pill", color=status["hex"])
        fw = comp.get("fit_width")
        if fw:          # width follows the label, right edge stays on the slot (estimated)
            paras = self.src.paras(ref, st)
            tw = sum(text_width_pt(t, fam, b_, i_, sz) for r in paras[0].runs for t, fam, b_, i_, sz in run_pieces(r)) / 72
            w = min(box[2], max(fw["min_w"], tw + 2 * fw["pad"]))
            box = [box[0] + box[2] - w, box[1], w, box[3]]
        sh = self.rect(slide, box, fill=status["fill"], shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=comp["corner_adj"][where])
        self.text_in(sh, self.src.paras(ref, st, align="ctr"), (0.05, 0, 0.05, 0), "ctr")

    def color_expr(self, expr, rule=None):
        """Box-style colour: 'none' | name | '<name>.fill' | '<box_map.line>' (taken from the box_map rule) → (hex, alpha)."""
        if expr in (None, "none"):
            return None, None
        if expr == "<box_map.line>":
            expr = (rule or {}).get("line", "rule_subtle")
        if str(expr).startswith("tag:"):
            return self.src.tag_hex(expr[4:]), None
        if re.fullmatch(r"[0-9A-Fa-f]{6}", str(expr)):
            return str(expr).upper(), None
        name, _, key = str(expr).partition(".")
        c = self.d.color(name)
        return str(c[key or "hex"]).upper(), (c.get("alpha") if not key else None)

    def box_of(self, sid):
        """Design box style for source shape `sid`: first deck.box_map rule whose fill/line match the shape
        (None = no fill / no line). Returns None when the plan has no rule for it."""
        sh = self.src.shapes.get(sid) or {}
        fill = ((sh.get("fill") or {}).get("hex") or "").upper() or None
        line = ((sh.get("line") or {}).get("hex") or "").upper() or None
        for i, rule in enumerate(self.plan["deck"].get("box_map", [])):
            m = rule["match"]
            mf, ml = (m.get("fill") or "").upper() or None, (m.get("line") or "").upper() or None
            if mf == fill and ml == line:
                bs = self.d.t["components"]["box_styles"][rule["style"]]
                fhex, falpha = self.color_expr(bs.get("fill"), rule)
                lhex, lalpha = self.color_expr(bs.get("line"), rule)
                self.box_used[i] = self.box_used.get(i, 0) + 1
                return {"name": rule["style"], "geom": bs.get("geom", "roundRect"), "corner_adj": bs.get("corner_adj"),
                        "fill": fhex, "fill_alpha": falpha, "line": lhex, "line_alpha": lalpha,
                        "line_pt": bs.get("line_pt", 0.86), "text": self.color_expr(bs.get("text"))[0] if bs.get("text") else None,
                        "dash": (sh.get("line") or {}).get("dash") not in (None, "solid"), "source": bs.get("source")}
        return None

    def paint_box(self, slide, box, bx):
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if bx["geom"] == "roundRect" else MSO_SHAPE.RECTANGLE
        return self.rect(slide, box, fill=bx["fill"], alpha=bx["fill_alpha"], line=bx["line"], width=bx["line_pt"],
                         dash="dash" if bx["dash"] else None, shape=shape,
                         adj=bx["corner_adj"] if shape == MSO_SHAPE.ROUNDED_RECTANGLE else None, line_alpha=bx["line_alpha"])

    def box_text(self, paras, bx):
        """Inside a status box the body-coloured runs take the box's text colour (R:S2–S8 pill rule)."""
        if not bx or not bx.get("text"):
            return paras
        white = self.d.hex("text")
        for p in paras:
            for r in p.runs:
                if r.style.color == white:
                    r.style = copy.copy(r.style)
                    r.style.color = bx["text"]
        return paras

    def cell(self, slide, box, sid, paras, insets, anchor, dashed=False, comp="panel"):
        """A content box painted with the box_map style of its source shape (fallback: the layout component)."""
        bx = self.box_of(sid) if sid else None
        sh = self.paint_box(slide, box, bx) if bx else self.panel(slide, box, dashed=dashed, comp=comp)
        self.text_in(sh, self.box_text(paras, bx), insets, anchor)
        return sh

    def panel(self, slide, box, dashed=False, comp="panel"):
        c = self.d.t["components"][comp]
        line = c.get("line", {})
        line_hex = self.d.hex(line.get("color")) if line.get("color") else None
        width = line.get("width_pt", 0.75)
        line_alpha = self.d.color(line["color"]).get("alpha") if line.get("color") else None
        fill = self.d.color(c["fill"])
        return self.rect(slide, box, fill=fill["hex"], alpha=fill.get("alpha"), line=line_hex, width=width,
                         dash="dash" if dashed else None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=c.get("corner_adj"),
                         line_alpha=line_alpha)

    def arrow(self, slide, x1, y, x2):
        c = self.d.t["components"]["arrow"]
        self.line(slide, x1, y, x2, y, self.d.hex(c["line"]["color"]), c["line"]["width_pt"], head=c.get("head"))

    def lead(self, slide, spec, box, style="lead_compact", page=None):
        ref = spec.get("slots", {}).get("lead")
        if not ref:
            return False
        st = self.d.style(style)
        paras = self.src.paras(ref, st)
        w = box[2] * 72
        size = st.size
        while size > self.d.min_size(st) and max_lines(self.restyle(paras, size / st.size), w) > 1:
            size -= 0.5
        if page is not None and size < st.size:
            page.fit.setdefault("shrunk", {})["lead"] = size
        out = self.restyle(paras, size / st.size)
        if max_lines(out, w) > 1 and page is not None:
            page.fit.setdefault("issues", []).append("lead 1줄 초과")
        self.text(self.slide_of(page), box, out)
        return True

    @staticmethod
    def slide_of(page):
        return page.slide

    def legend_paras(self, items, st, tag_size=None):
        runs = []
        for k, it in enumerate(items):
            tag = it["tag"]["t"] if isinstance(it["tag"], dict) else it["tag"]
            tst = self.d.style("tag", color=self.src.tag_hex(tag))
            if tag_size:
                tst.size = tag_size
            runs.append(Run(tag + " ", tst))
            runs.append(Run(self.src.text(it["label"]) + ("      " if k < len(items) - 1 else ""), st))
        return [Para(runs)]


# ---------------------------------------------------------------- theme fonts

def replace_theme_fonts(prs, design: Design):
    """Theme ea / Hangul script fonts → Pretendard (and 맑은 고딕 latin → Arial / Roboto Light), so runs
    that miss an explicit <a:ea> do not fall back to 맑은 고딕 (DECISIONS #9, #17)."""
    heading, body = design.t["fonts"]["heading"], design.t["fonts"]["body"]
    want = {"majorFont": (heading["latin"], heading["ea"]), "minorFont": (body["latin"], body["ea"])}
    changes, seen = [], set()
    for m in prs.slide_masters:
        part = m.part.part_related_by(RT.THEME)
        if id(part) in seen:
            continue
        seen.add(id(part))
        root = etree.fromstring(part.blob)
        for which, (latin, ea) in want.items():
            f = root.find(f".//{qa('fontScheme')}/{qa(which)}")
            if f is None:
                continue
            la = f.find(qa("latin"))
            if la is not None and la.get("typeface") in ("맑은 고딕", "Malgun Gothic", ""):
                changes.append(f"{part.partname} {which} latin {la.get('typeface')!r} → {latin}")
                la.set("typeface", latin)
                la.attrib.pop("panose", None)
            e = f.find(qa("ea"))
            if e is not None and e.get("typeface") != ea:
                changes.append(f"{part.partname} {which} ea {e.get('typeface')!r} → {ea}")
                e.set("typeface", ea)
            for sf in f.findall(qa("font")):
                if sf.get("script") == "Hang" and sf.get("typeface") != ea:
                    changes.append(f"{part.partname} {which} Hang {sf.get('typeface')!r} → {ea}")
                    sf.set("typeface", ea)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return changes


# ---------------------------------------------------------------- layouts

def cell_paras_io(b: Builder, cell: dict, which: str, ratio: float, mins: dict):
    d, src = b.d, b.src
    st_line, st_note = d.style("io_line"), d.style("io_note")
    paras = []
    if which == "step":
        st_no, st_head = d.style("io_step_no"), d.style("io_headline")
        runs = []
        if cell.get("no"):
            runs.append(Run(src.text(cell["no"]) + "  ", st_no))
        runs.append(Run(src.text(cell["headline"]), st_head))
        for t in cell.get("tags", []):
            tag = src.text(t)
            runs.append(Run("  " + tag, d.style("tag", color=src.tag_hex(tag))))
        paras.append(Para(runs))
        for r in cell.get("lines", []):
            paras += src.paras(r, st_line, keep_runs=True)
        for r in cell.get("note", []):
            paras += src.paras(r, st_note, keep_runs=True)
    else:
        for r in cell.get("lines", []):
            paras += src.paras(r, st_line, keep_runs=True)
        for r in cell.get("secondary", []):
            paras += src.paras(r, st_note, keep_runs=True)
    return b.restyle(paras, ratio, mins)


def render_io_steps(b: Builder, spec: dict):
    d = b.d
    lay = d.t["layouts"]["io_steps"]
    g = Geo(d, lay)
    reg = spec["regions"]
    rows = reg["rows"]
    comp = d.t["components"]["step_cell"]
    ins = comp["insets"]
    has_foot = bool(reg.get("footnote"))
    R = lay["rows"]
    top, bottom = R["top"], R["bottom"] if has_foot else R["bottom_without_footnote"]
    cols = lay["columns"]
    styles = ["io_line", "io_note", "io_step_no", "io_headline", "tag"]
    mins = {s: d.min_size(d.style(s)) for s in styles}

    def fits(group, ratio):
        n = len(group)
        rh = (bottom - top - R["gap"] * (n - 1)) / n
        worst = []
        for row in group:
            for which in ("in", "step", "out"):
                c = cols[which]
                w = (g.w(c["w"]) - ins[0] - ins[2]) * 72
                hh = (g.h(rh) - ins[1] - ins[3]) * 72
                need = text_height(cell_paras_io(b, row[which], which, ratio, mins), w)
                if need > hh + 0.5:
                    worst.append((which, round(need - hh, 1)))
        return worst

    # shrink → split (by rows): smallest ratio ≥ min that fits; else split in two and retry
    def plan_groups(group):
        ratio = 1.0
        while ratio >= 0.85 - 1e-9:
            if not fits(group, ratio):
                return [(group, ratio)]
            ratio = round(ratio - 0.025, 3)
        if len(group) <= 1:
            return [(group, 0.85)]
        half = math.ceil(len(group) / 2)
        return plan_groups(group[:half]) + plan_groups(group[half:])

    groups = plan_groups(rows)
    for gi, (group, ratio) in enumerate(groups):
        part = (gi + 1, len(groups)) if len(groups) > 1 else None
        page = b.new_slide(spec, part)
        slide = page.slide
        if len(groups) > 1:
            page.fit["split"] = f"{len(rows)}행 → {[len(x) for x, _ in groups]}"
        if ratio < 1:
            page.fit.setdefault("shrunk", {})["cells"] = f"×{ratio:.3f}"
        left = fits(group, ratio)
        if left:
            page.fit.setdefault("issues", []).append(f"셀 넘침 {left}")
        b.lead(slide, spec, g.box(lay["regions"]["lead"]["box"]), page=page)
        hst = d.style(lay["header"]["style"])
        for which, ref in zip(("in", "step", "out"), reg.get("header", [])):
            c = cols[which]
            b.text(slide, [g.x(c["x"]), g.y(lay["header"]["y"]), g.w(c["w"]), g.h(lay["header"]["h"])], b.src.paras(ref, hst))
        n = len(group)
        rh = (bottom - top - R["gap"] * (n - 1)) / n
        for i, row in enumerate(group):
            y = top + i * (rh + R["gap"])
            for which in ("in", "step", "out"):
                c = cols[which]
                box = [g.x(c["x"]), g.y(y), g.w(c["w"]), g.h(rh)]
                cell = row[which]
                first = cell.get("headline") or next(iter(cell.get("lines", []) + cell.get("secondary", [])), None)
                b.cell(slide, box, first and first["src"], cell_paras_io(b, cell, which, ratio, mins), ins,
                       comp.get("anchor", "ctr"), dashed=which == "step" and cell.get("dashed"), comp="step_cell")
            ymid = g.y(y + rh / 2)
            gap = lay["arrows"]["gap"]
            for a_, z_ in lay["arrows"]["between"]:
                x1 = g.x(cols[a_]["x"] + cols[a_]["w"]) + 0.04
                x2 = g.x(cols[z_]["x"]) - 0.04
                if x2 - x1 > 0.05 and gap > 0:
                    b.arrow(slide, x1, ymid, x2)
        if has_foot:
            st = d.style("io_note")
            fb = g.box(lay["regions"]["footnote"]["box"])
            paras = b.src.paras(reg["footnote"], st, keep_runs=True)
            if max_lines(paras, fb[2] * 72) > 1:
                page.fit.setdefault("issues", []).append("footnote 1줄 초과")
            b.text(slide, fb, paras)


def render_card_grid(b: Builder, spec: dict):
    d = b.d
    lay = d.t["layouts"]["card_grid"]
    var = spec.get("layout_variant")
    vspec = (lay.get("variants") or {}).get(var, {}) if var else {}
    g = Geo(d, lay)
    reg = spec["regions"]
    page = b.new_slide(spec)
    slide = page.slide
    cols = vspec.get("columns", lay["columns"])
    card = dict(lay["card"])
    card.update(vspec.get("card", {}))
    ins = lay["card"]["insets"]
    shift = 0.0
    if spec.get("slots", {}).get("lead"):           # card_grid has no lead region: put it on top, push the grid down (estimated)
        ca = d.t["chrome"]["collision"]["content_area"]
        b.lead(slide, spec, g.box([ca[0], ca[1], ca[2], 0.4]), page=page)
        shift = 0.6
        page.fit.setdefault("reflow", []).append("lead 위 배치, 카드 0.6″ 아래로")
    banner = reg.get("banner")
    if var == "stats":
        rows = vspec["rows"]
    else:
        rows = lay["rows"]["with_banner" if banner else "without_banner"]
    if banner:
        bb = g.box(lay["regions"]["banner"]["box"])
        bb[1] += shift
        hst, bst = d.style("card_heading"), d.style("card_body")
        runs = []
        for p in b.src.paras(banner["heading"], hst, keep_runs=True):
            runs += p.runs
        runs.append(Run("   ", bst))
        for r in banner.get("body", []):
            for p in b.src.paras(r, bst, keep_runs=True):
                runs += p.runs
        paras = [Para(runs)]
        bins = (0.25, 0.08, 0.25, 0.08)
        w = (bb[2] - bins[0] - bins[2]) * 72
        ratio = 1.0
        while ratio > 0.85 and text_height(b.restyle(paras, ratio), w) > (bb[3] - bins[1] - bins[3]) * 72:
            ratio -= 0.025
        if ratio < 1:
            page.fit.setdefault("shrunk", {})["banner"] = f"×{ratio:.3f}"
        b.cell(slide, bb, banner["heading"]["src"], b.restyle(paras, ratio), bins, "ctr", dashed=banner.get("dashed"))
    cards = reg.get("cards", [])
    ncol = cols["count"]
    for i, c in enumerate(cards):
        r, k = divmod(i, ncol)
        box = g.box([cols["x"][k], rows["y"][r] + shift, cols["w"], rows["h"]])
        w = (box[2] - ins[0] - ins[2]) * 72
        if var == "stats":
            nst, lst, dst = d.style(card["number"]["style"]), d.style(card["label"]["style"]), d.style(card["detail"]["style"])
            paras = [Para([Run(b.src.text(c["number"]), nst), Run("  ", lst), Run(b.src.text(c["label"]), lst)])]
            for r_ in c.get("detail", []):
                paras += b.src.paras(r_, dst, keep_runs=True)
            if paras[1:]:
                paras[1].space_before = 4
        else:
            hst, bst = d.style(card["heading"]["style"]), d.style(card["body"]["style"])
            paras = b.src.paras(c["heading"], hst)          # heading keeps the design colour (source blue is decoration)
            for j, r_ in enumerate(c.get("body", [])):
                ps = b.src.paras(r_, bst, keep_runs=True)
                if j == 0 and ps:
                    ps[0].space_before = card["body"].get("gap", 0.08) * 72
                paras += ps
        ratio = 1.0
        mins = {s: d.min_size(d.style(s)) for s in ("card_body", "card_heading", "io_note", "stat_number", "stat_label")}
        hh = (box[3] - ins[1] - ins[3]) * 72
        while ratio > 0.85 and text_height(b.restyle(paras, ratio, mins), w) > hh:
            ratio = round(ratio - 0.025, 3)
        out = b.restyle(paras, ratio, mins)
        if text_height(out, w) > hh + 0.5:
            page.fit.setdefault("issues", []).append(f"카드 {i + 1} 넘침")
        if ratio < 1:
            page.fit.setdefault("shrunk", {})[f"card{i + 1}"] = f"×{ratio:.3f}"
        sid = (c.get("heading") or c.get("number"))["src"]
        b.cell(slide, box, sid, out, ins, "t" if var != "stats" else "ctr")
    if reg.get("legend"):
        lb = g.box(lay["regions"]["legend"]["box"])
        b.text(slide, lb, b.legend_paras(reg["legend"], d.style("legend")))


def render_cover(b: Builder, spec: dict, closing=False):
    d = b.d
    lay = d.t["layouts"]["closing" if closing else "cover"]
    page = b.new_slide(spec, chrome=False)
    slide = page.slide
    for bg in lay.get("background", []):
        b.rect(slide, bg["box"], fill=d.hex(bg["fill"]))
    slots, reg = spec.get("slots", {}), spec.get("regions", {})
    y_next = None
    for name, r in lay["regions"].items():
        if r.get("kind") == "picture":
            b.picture(slide, b.asset(r["asset"]), r["box"])
        elif r.get("kind") == "line":
            bx = r["box"]
            b.line(slide, bx[0], bx[1], bx[0] + bx[2], bx[1] + bx[3], d.hex(r["color"]), r["width_pt"])
        elif r.get("component") == "device_art":
            for piece in d.t["components"]["device_art"]["pieces"]:
                b.picture(slide, b.asset(piece["asset"]), piece["box"])
        elif name in ("flow", "legend"):
            continue
        elif slots.get(name) or reg.get(name):
            ref = slots.get(name) or reg.get(name)
            tdef = d.t["type"][r["style"]]
            st = d.style(r["style"])
            ins = r.get("insets", [0, 0, 0, 0])
            pk = {"line_spacing": float(tdef.get("line_spacing", 1.0))}
            if tdef.get("bullet"):
                pk.update(bullet=tdef["bullet"], bullet_color=st.color, marL=tdef.get("marL", 0), indent=tdef.get("indent", 0))
            paras = b.src.paras(ref, st, **pk)
            box = list(r["box"])
            w = (box[2] - ins[0] - ins[2]) * 72
            ml = r.get("max_lines")
            ratio = 1.0
            while ratio > 0.85 and ml and max(wrap_count(p, w) for p in b.restyle(paras, ratio)) > ml:
                ratio = round(ratio - 0.025, 3)
            out = b.restyle(paras, ratio)
            if ml and max(wrap_count(p, w) for p in out) > ml:
                page.fit.setdefault("issues", []).append(f"{name} {ml}줄 초과")
            if ratio < 1:
                page.fit.setdefault("shrunk", {})[name] = f"×{ratio:.3f}"
            if r.get("autofit") == "shape":         # spAutoFit: the box takes the text height
                box[3] = max(box[3], text_height(out, w) / 72 + ins[1] + ins[3])
            b.text(slide, box, out, ins, "t", r.get("autofit", "none"))
            y_next = max(y_next or 0, box[1] + box[3])
    if not closing and reg.get("flow"):
        r = lay["regions"]["flow"]
        it = r["item"]
        items = reg["flow"][: it["max_items"]]
        bx = list(r["box"])
        if y_next and y_next > bx[1] - 0.1:
            page.fit.setdefault("issues", []).append(f"meta 아래 끝 {y_next:.2f}″ — flow({bx[1]}″)와 겹침")
        n = len(items)
        w = (bx[2] - it["gap"] * (n - 1)) / n
        st = d.style(it["style"])
        comp = d.t["components"][it["component"]]
        for i, item in enumerate(items):
            x = bx[0] + i * (w + it["gap"])
            b.cell(slide, [x, bx[1], w, bx[3]], item["text"]["src"], b.src.paras(item["text"], st, keep_runs=True, align="ctr"),
                   comp["insets"], "ctr", dashed=item.get("dashed"), comp=it["component"])
            if i < n - 1:
                b.arrow(slide, x + w + 0.06, bx[1] + bx[3] / 2, x + w + it["gap"] - 0.06)
    if not closing and reg.get("legend"):
        r = lay["regions"]["legend"]
        b.text(slide, r["box"], b.legend_paras(reg["legend"], d.style(r["style"])))


def render_image_full(b: Builder, spec: dict):
    lay = b.d.t["layouts"]["image_full"]
    page = b.new_slide(spec, chrome=False)
    img = spec["regions"]["image"]
    path = b.base_dir / img["file"]
    crop = img.get("crop_in")
    if crop:
        sw = b.content["slide_size"]["w_in"]
        with Image.open(path) as im:
            k = im.size[0] / sw
            box = tuple(int(round(v * k)) for v in (crop[0], crop[1], crop[0] + crop[2], crop[1] + crop[3]))
            out = b.build_dir / f"crop-{spec['out']:02d}.png"
            im.crop(box).save(out)
            path = out
    b.picture(page.slide, path, lay["regions"]["image"]["box"], fit=img.get("fit", "contain"),
              align=lay["regions"]["image"].get("align", "center"))


_src_prs: dict = {}
FILL_TAGS = ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
DROP_TAGS = ("effectLst", "effectDag", "scene3d", "sp3d")


def source_slide(path, n):
    if path not in _src_prs:
        _src_prs[path] = Presentation(str(path))
    return _src_prs[path].slides[n - 1]


def cnv_id(el):
    c = el.find(".//{%s}cNvPr" % P)
    return int(c.get("id")) if c is not None else None


def line_hex(b: Builder, hexv):
    """Line colour for a redrawn diagram: deck.line_map, then deck.color_map, else rule_subtle."""
    lm = {k.upper(): v for k, v in (b.plan["deck"].get("line_map") or {}).items()}
    v = lm.get((hexv or "").upper())
    if v:
        return b.src.tag_hex(v[4:]) if v.startswith("tag:") else b.d.hex(v)
    return b.src.mapped(hexv) or b.d.hex(b.d.t["layouts"]["diagram_restyle"]["shape"]["line_default"])


def diagram_style(b: Builder, run: dict, scale: float, min_size: float, box_text=None):
    d = b.d
    t = d.t["layouts"]["diagram_restyle"]["text"]
    role = t["bold" if run.get("bold") else "regular"]
    role = d.font_over.get(role, role)
    f = d.t["fonts"][role]
    hexv = b.src.mapped((run.get("color") or {}).get("hex")) or d.hex("text")
    if box_text and hexv == d.hex("text"):
        hexv = box_text
    return Style(latin=f["latin"], ea=f["ea"], size=max(min_size, round((run.get("size") or 18) * scale * 2) / 2),
                 bold=bool(run.get("bold")), italic=bool(run.get("italic")), color=hexv, name="diagram")


def restyle_element(b: Builder, el, cs: dict, scale: float, min_size: float, page):
    """Repaint one cloned sp/cxnSp: panel fill, mapped line, no effects, design fonts. Geometry untouched."""
    d = b.d
    st = el.find(f"{{{P}}}style")
    if st is not None:
        el.remove(st)
    bx = None
    sppr = el.find(f"{{{P}}}spPr")
    if sppr is not None:
        for ch in list(sppr):
            if ch.tag in (qa(t) for t in FILL_TAGS + DROP_TAGS):
                sppr.remove(ch)
        geom = sppr.find(qa("prstGeom"))
        if geom is None:
            geom = sppr.find(qa("custGeom"))
        at = list(sppr).index(geom) + 1 if geom is not None else len(sppr)
        bx = b.box_of(cs["id"]) if cs["kind"] == "shape" else None
        if bx:
            if bx["fill"]:
                sf = etree.Element(qa("solidFill"))
                clr = etree.SubElement(sf, qa("srgbClr"))
                clr.set("val", bx["fill"])
                if (bx["fill_alpha"] or 1) < 1:
                    etree.SubElement(clr, qa("alpha")).set("val", str(int(bx["fill_alpha"] * 100000)))
            else:
                sf = etree.Element(qa("noFill"))
        elif cs["kind"] == "shape" and cs.get("fill"):
            fill = d.color(d.t["layouts"]["diagram_restyle"]["shape"]["fill"])
            sf = etree.Element(qa("solidFill"))
            clr = etree.SubElement(sf, qa("srgbClr"))
            clr.set("val", fill["hex"])
            if fill.get("alpha", 1) < 1:
                etree.SubElement(clr, qa("alpha")).set("val", str(int(fill["alpha"] * 100000)))
        else:
            sf = etree.Element(qa("noFill"))
        if el.tag != f"{{{P}}}cxnSp":
            sppr.insert(at, sf)
        ln = sppr.find(qa("ln"))
        if bx and bx["line"]:
            if ln is not None:
                sppr.remove(ln)
            ln = etree.SubElement(sppr, qa("ln"))
            ln.set("w", str(int(bx["line_pt"] * 12700)))
            lf = etree.SubElement(ln, qa("solidFill"))
            etree.SubElement(lf, qa("srgbClr")).set("val", bx["line"])
            if bx["dash"]:
                etree.SubElement(ln, qa("prstDash")).set("val", "dash")
        elif cs.get("line"):
            if ln is None:
                ln = etree.SubElement(sppr, qa("ln"))
                ln.set("w", str(int(cs["line"].get("width_pt", 0.75) * 12700)))
                if cs["line"].get("dash") not in (None, "solid"):
                    etree.SubElement(ln, qa("prstDash")).set("val", cs["line"]["dash"])
            for ch in list(ln):
                if ch.tag in (qa(t) for t in FILL_TAGS):
                    ln.remove(ch)
            lf = etree.Element(qa("solidFill"))
            etree.SubElement(lf, qa("srgbClr")).set("val", line_hex(b, cs["line"]["hex"]))
            ln.insert(0, lf)
        else:
            if ln is not None:
                sppr.remove(ln)
            ln = etree.SubElement(sppr, qa("ln"))
            etree.SubElement(ln, qa("noFill"))
        etree.SubElement(sppr, qa("effectLst"))
    tx = el.find(f"{{{P}}}txBody")
    if tx is None or not cs.get("text"):
        return
    bp = tx.find(qa("bodyPr"))
    for k, dflt in (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720)):
        bp.set(k, str(int(int(bp.get(k, dflt)) * scale)))
    paras_xml = tx.findall(qa("p"))
    measured = []
    for px, pc in zip(paras_xml, cs["text"]["paragraphs"]):
        runs_xml = px.findall(qa("r"))
        runs = []
        for k, rx in enumerate(runs_xml):
            rc = pc["runs"][k] if k < len(pc["runs"]) else {}
            s2 = diagram_style(b, rc, scale, min_size, box_text=(bx or {}).get("text"))
            run = Run(rx.findtext(qa("t")) or "", s2)
            runs.append(run)
            parent, at = rx.getparent(), list(rx.getparent()).index(rx)
            parent.remove(rx)
            for j, (text, st_) in enumerate(coverage_runs(run)):     # same glyph-coverage split as write_paras
                nr = etree.Element(qa("r"))
                rpr(nr, st_)
                etree.SubElement(nr, qa("t")).text = text
                parent.insert(at + j, nr)
        end = px.find(qa("endParaRPr"))
        if end is not None and runs:
            end.set("sz", str(int(runs[-1].style.size * 100)))
        if runs:
            measured.append(Para(runs, align=pc.get("align") or "l"))
    box = cs["box"]
    w = box[2] / EMU_IN * scale * 72 - (int(bp.get("lIns")) + int(bp.get("rIns"))) / EMU_IN * 72
    h = box[3] / EMU_IN * scale * 72 - (int(bp.get("tIns")) + int(bp.get("bIns"))) / EMU_IN * 72
    if measured and w > 0:
        need = text_height(measured, w)
        if need > h + 1 and cs["kind"] == "shape":
            page.fit.setdefault("issues", []).append(f"{cs['id']} 글자 넘침 {need - h:.0f}pt")


def render_diagram_restyle(b: Builder, spec: dict):
    d = b.d
    lay = d.t["layouts"]["diagram_restyle"]
    g = Geo(d, lay)
    reg = spec.get("regions", {})
    dia = reg.get("diagram", {})
    exclude = set(dia.get("exclude", []))
    cslide = b.content["slides"][spec["src"] - 1]
    by_xml = {s["xml_id"]: s for s in cslide["shapes"] if not s.get("group")}
    keep = [s for s in by_xml.values() if s["id"] not in exclude]
    x0 = min(s["box"][0] for s in keep) / EMU_IN
    y0 = min(s["box"][1] for s in keep) / EMU_IN
    x1 = max(s["box"][0] + s["box"][2] for s in keep) / EMU_IN
    y1 = max(s["box"][1] + s["box"][3] for s in keep) / EMU_IN
    rb = g.box(lay["regions"]["diagram"]["box"])
    if reg.get("footnote"):
        rb[3] = g.y(lay["regions"]["diagram"]["bottom_with_footnote"]) - rb[1]
    if not spec.get("slots", {}).get("lead"):
        top = g.y(lay["regions"]["lead"]["box"][1])
        rb[3] += rb[1] - top
        rb[1] = top
    scale = min(rb[2] / (x1 - x0), rb[3] / (y1 - y0))
    ox, oy = rb[0] + (rb[2] - (x1 - x0) * scale) / 2, rb[1]
    page = b.new_slide(spec)
    slide = page.slide
    page.fit["scale"] = round(scale, 3)
    smin = d.t["overflow"]["per_layout"].get("diagram_restyle", {}).get("scale_min", 0.7)
    if scale < smin:
        page.fit.setdefault("issues", []).append(f"배율 {scale:.2f} < {smin} — 글자가 작음")
    b.lead(slide, spec, g.box(lay["regions"]["lead"]["box"]), page=page)

    src_path = Path(b.plan["source"])
    src_tree = source_slide(src_path if src_path.exists() else b.root / src_path, spec["src"]).shapes._spTree
    tree = slide.shapes._spTree
    next_id = max(int(c.get("id")) for c in tree.iter(f"{{{P}}}cNvPr")) + 1
    remap, clones = {}, []
    min_size = float(lay["text"].get("min_size", 8))
    for el in src_tree.iterchildren():
        xid = cnv_id(el) if el.tag != f"{{{P}}}nvGrpSpPr" else None
        if xid is None or xid not in by_xml or by_xml[xid]["id"] in exclude:
            continue
        if el.tag == f"{{{P}}}pic" or el.find(f".//{{{P}}}pic") is not None or el.find(f".//{qa('graphicData')}") is not None:
            page.fit.setdefault("issues", []).append(f"{by_xml[xid]['id']} 그림/표/차트는 다시 칠하지 않음 — 빠짐")
            continue
        new = copy.deepcopy(el)
        xfrm = new.find(f"{{{P}}}spPr/{qa('xfrm')}")
        if xfrm is None:
            xfrm = new.find(f"{{{P}}}grpSpPr/{qa('xfrm')}")
        off, ext = xfrm.find(qa("off")), xfrm.find(qa("ext"))
        off.set("x", str(int(round((ox + (int(off.get("x")) / EMU_IN - x0) * scale) * EMU_IN))))
        off.set("y", str(int(round((oy + (int(off.get("y")) / EMU_IN - y0) * scale) * EMU_IN))))
        ext.set("cx", str(int(round(int(ext.get("cx")) * scale))))
        ext.set("cy", str(int(round(int(ext.get("cy")) * scale))))
        # every sp/cxnSp inside (a group's children keep child coordinates — only the group frame moves)
        targets = [new] if new.tag != f"{{{P}}}grpSp" else list(new.iter(f"{{{P}}}sp", f"{{{P}}}cxnSp"))
        for t_el in targets:
            cs = by_xml.get(cnv_id(t_el)) or next((s for s in cslide["shapes"] if s["xml_id"] == cnv_id(t_el)), None)
            if cs:
                restyle_element(b, t_el, cs, scale, min_size, page)
        for c in new.iter(f"{{{P}}}cNvPr"):
            remap[int(c.get("id"))] = next_id
            c.set("id", str(next_id))
            next_id += 1
        clones.append(new)
    for new in clones:
        for tag in ("stCxn", "endCxn"):
            for c in new.iter(qa(tag)):
                if int(c.get("id")) in remap:
                    c.set("id", str(remap[int(c.get("id"))]))
                else:
                    c.getparent().remove(c)
        tree.append(new)
    if reg.get("footnote"):
        st = d.style(lay["regions"]["footnote"]["style"])
        fb = g.box(lay["regions"]["footnote"]["box"])
        paras = b.src.paras(reg["footnote"], st, keep_runs=True)
        if max_lines(paras, fb[2] * 72) > 1:
            page.fit.setdefault("issues", []).append("footnote 1줄 초과")
        b.text(slide, fb, paras)


def render_generic(b: Builder, spec: dict):
    """Layouts whose regions are plain text boxes (section, references, bullets…): each region with a
    `box` gets its referenced text in the region style. Components other than bullets are not drawn."""
    d = b.d
    lay = d.t["layouts"][spec["layout"]]
    g = Geo(d, lay)
    page = b.new_slide(spec)
    slide = page.slide
    for name, r in (lay.get("regions") or {}).items():
        val = spec.get("regions", {}).get(name) or spec.get("slots", {}).get(name)
        if not val or "box" not in r:
            continue
        refs = val if isinstance(val, list) else [val]
        box = g.box(r["box"])
        if r.get("component") == "bullets":
            bc = d.t["components"]["bullets"]
            paras = []
            for ref in refs:
                lvl = "l2" if ref.get("level", 0) else "l1"
                spec_l = bc[lvl]
                st = d.style(spec_l["style"])
                paras += b.src.paras(ref, st, bullet=bc["char"], bullet_color=d.hex(bc["color"]),
                                     marL=spec_l["marL"], indent=spec_l["indent"], space_before=spec_l["space_before"])
        elif r.get("style"):
            st = d.style(r["style"])
            paras = []
            for ref in refs:
                paras += b.src.paras(ref, st)
        else:
            if r.get("fit"):         # image region
                img = val.get("file") if isinstance(val, dict) else None
                if img:
                    b.picture(slide, b.base_dir / img, box, fit=r["fit"], align=r.get("align", "center"))
            continue
        if text_height(paras, box[2] * 72) > box[3] * 72 + 0.5:
            page.fit.setdefault("issues", []).append(f"{name} 넘침")
        b.text(slide, box, paras)


RENDERERS = {
    "io_steps": render_io_steps,
    "card_grid": render_card_grid,
    "cover": render_cover,
    "closing": lambda b, s: render_cover(b, s, closing=True),
    "image_full": render_image_full,
    "diagram_restyle": render_diagram_restyle,
    "section": render_generic, "references": render_generic, "bullets": render_generic,
    "bullets_image": render_generic, "screenshot_steps": render_generic,
}
NOT_YET = ("table", "toc", "intro_columns", "visual_kv", "image_pair_kv")


# ---------------------------------------------------------------- report

def write_report(path: Path, b: Builder, plan, tokens, variant, out_pptx: Path, fonts_note: list[str] | None):
    d = b.d
    lines = [f"# 변환 보고서 — {Path(plan['source']).name}", "",
             f"- 출력: `{out_pptx}` ({len(b.pages)}장, 원본 {len(b.content['slides'])}장)",
             f"- 변형: `{variant}`" + (" (기본)" if variant == tokens["meta"]["default_chrome"] else ""),
             f"- 디자인: `{plan.get('design', 'DESIGN.md')}` {tokens['meta']['name']} {tokens['meta']['version']}", ""]
    lines += ["## 슬라이드", "", "| 출력 | 원본 | 레이아웃 | 실측/추정 | 맞춤 | 수작업 |", "|---|---|---|---|---|---|"]
    used = {}
    for pg in b.pages:
        sp = pg.spec
        lay = sp["layout"] + (f".{sp['layout_variant']}" if sp.get("layout_variant") else "")
        prov = provenance(tokens["layouts"][sp["layout"]], variant)
        if sp.get("layout_variant"):
            vs = tokens["layouts"][sp["layout"]].get("variants", {}).get(sp["layout_variant"], {})
            if vs.get("source") == "estimated":
                prov = "estimated"
        used.setdefault(lay, prov)
        fit = []
        if pg.fit.get("split"):
            fit.append(f"분할 {pg.fit['split']}")
        if pg.part:
            fit.append(f"{pg.part[0]}/{pg.part[1]}")
        for k, v in (pg.fit.get("shrunk") or {}).items():
            fit.append(f"축소 {k} {v}")
        if pg.fit.get("scale"):
            fit.append(f"배율 ×{pg.fit['scale']}")
        for v in pg.fit.get("reflow", []):
            fit.append(f"재배치 {v}")
        for v in pg.fit.get("issues", []):
            fit.append(f"**넘침** {v}")
        lines.append(f"| {pg.out_n} | S{sp['src']} | `{lay}` | {'실측' if prov == 'measured' else '추정'} | {'; '.join(fit) or '-'} | {sp.get('manual') or ''} |")
    est = [k for k, v in used.items() if v == "estimated"]
    lines += ["", "## 추정으로 그린 레이아웃", ""]
    lines += [f"- `{k}`" for k in est] or ["- 없음"]
    lines += ["", "## 요약·분할", ""]
    splits = [pg for pg in b.pages if pg.fit.get("split")]
    plan_splits = sorted({(s["src"], s["split"]["of"]) for s in plan["slides"] if s.get("split")})
    if plan_splits:
        lines += [f"- 계획 분할: 원본 S{s} → {m}장 (plan.json)" for s, m in plan_splits]
    seen = set()
    for pg in splits:
        if pg.spec["out"] not in seen:
            seen.add(pg.spec["out"])
            lines.append(f"- 넘침 분할: 원본 S{pg.spec['src']} {pg.fit['split']} (축소 하한에서도 셀이 넘침)")
    summ = [s for s in plan["slides"] if "summarize" in (s.get("overflow") or {}).get("plan", [])]
    lines += [f"- 요약: 원본 S{s['src']} — 원문은 노트에 보존" for s in summ]
    if not (plan_splits or splits or summ):
        lines.append("- 없음")
    lines.append("- 원문 텍스트는 바꾸지 않았다(모든 텍스트는 plan.json의 원본 shape 참조, `check_plan.py` 검사).")
    lines += ["", "## 수작업 필요", ""]
    man = [(pg.out_n, pg.spec) for pg in b.pages if pg.spec.get("manual")]
    lines += [f"- 출력 {n} (원본 S{s['src']}): {s['manual']}" for n, s in man] or ["- 없음"]
    issues = [(pg.out_n, i) for pg in b.pages for i in pg.fit.get("issues", [])]
    if issues:
        lines += ["", "## 남은 넘침 (확인 필요)", ""] + [f"- 출력 {n}: {i}" for n, i in issues]
    lines += ["", "## 빠진 원본 텍스트", ""]
    drops = [(s["src"], dr) for s in plan["slides"] for dr in s.get("dropped", [])]
    lines += [f"- S{n} `{dr['src']}`: {dr['why']}" for n, dr in drops] or ["- 없음"]
    lines += ["", "## 상자 스타일 대응 (`deck.box_map` → DESIGN.md `components.box_styles`)", ""]
    bs = tokens["components"].get("box_styles", {})
    for i, rule in enumerate(plan["deck"].get("box_map", [])):
        m = rule["match"]
        key = rule["style"] + (f"({rule['line']})" if rule.get("line") else "")
        src_ = bs.get(rule["style"], {}).get("source", "?")
        lines.append(f"- 채움 {m.get('fill') or '없음'} / 선 {m.get('line') or '없음'} → `{key}` "
                     f"{'실측 ' + bs[rule['style']].get('evidence', '') if src_ == 'measured' else '추정'} · {b.box_used.get(i, 0)}곳")
    lines.append("- 원본 상자가 점선이면 대응 스타일에 점선만 유지(추정). 대응이 없는 상자는 레이아웃 기본(panel/step_cell).")
    lines += ["", "## 서식 대응 (추정)", "",
              "- 원본 런 색은 `deck.color_map`으로 디자인 색에 대응시켰다. 고정폭(코드) 런은 글꼴을 본문 글꼴로 바꾸고 색만 유지한다.",
              "- 글자 태그 색: " + ", ".join(f"{k} → {v['color']}" + (f" ({v['note']})" if v.get("note") else "") for k, v in plan["deck"].get("tags", {}).items()),
              "", "## 폰트", ""]
    lines += [f"- 테마 폰트 교체: {c}" for c in b.theme_changes] or ["- 테마 폰트 교체 없음"]
    if missing_fonts:
        lines += [f"- 측정용 폰트 대체: {m}" for m in sorted(missing_fonts)]
    if fonts_note is not None:
        lines += fonts_note
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def render_and_check(out_pptx: Path, png_dir: Path):
    sh = Path(__file__).resolve().parent / "render.sh"
    res = subprocess.run([str(sh), str(out_pptx), str(png_dir)], capture_output=True, text=True)
    if res.returncode:
        return [f"- 렌더 실패: {res.stderr.strip()[:300]}"]
    fonts = [l.strip() for l in res.stdout.splitlines() if l.startswith("  ")]
    families = ("Arial", "Roboto", "Pretendard")        # faces DESIGN.md asks for; anything else is a substitute
    bad = [f for f in fonts if "fallback" in f or not f.split()[0].startswith(families)]
    note = [f"- 렌더(`render.sh`) 임베드 폰트: {', '.join(f.split()[0] for f in fonts)}"]
    note.append("- 폰트 대체: " + ("**있음** — " + ", ".join(bad) if bad else "없음"))
    return note


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan", type=Path)
    ap.add_argument("--variant", help="chrome variant (default: plan.variant or DESIGN.md default)")
    ap.add_argument("--design", type=Path, help="DESIGN.md (default: plan.design relative to cwd)")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--render", action="store_true", help="render with render.sh into work/x/build/png and report font substitution")
    a = ap.parse_args()

    plan = json.loads(a.plan.read_text(encoding="utf-8"))
    base_dir = a.plan.resolve().parent
    content = json.loads((base_dir / plan.get("content", "content.json")).read_text(encoding="utf-8"))
    tokens = load_tokens(a.design or Path(plan.get("design", "DESIGN.md")))
    variant = a.variant or plan.get("variant") or tokens["meta"]["default_chrome"]
    if variant not in tokens["chrome"]:
        sys.exit(f"unknown variant {variant!r}; choose from {list(tokens['chrome'])}")
    stem = Path(plan["source"]).stem
    out = a.out or Path("outputs") / (f"{stem}.pptx" if variant == tokens["meta"]["default_chrome"] else f"{stem}.{variant}.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    build_dir = base_dir / "build" / variant
    build_dir.mkdir(parents=True, exist_ok=True)

    design = Design(tokens, variant)
    src = Source(content, plan, design)
    b = Builder(design, plan, content, src, base_dir, build_dir)
    todo = [s["layout"] for s in plan["slides"] if s["layout"] in NOT_YET or s["layout"] not in RENDERERS]
    if todo:
        sys.exit(f"layouts not implemented in build_deck yet: {sorted(set(todo))}")
    for spec in plan["slides"]:
        RENDERERS[spec["layout"]](b, spec)
    b.prs.save(str(out))

    fit = [{"out": pg.out_n, "src": pg.spec["src"], "plan_out": pg.spec["out"], "layout": pg.spec["layout"],
            "part": pg.part, **pg.fit} for pg in b.pages]
    (build_dir / "build.json").write_text(json.dumps(fit, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    fonts_note = render_and_check(out, build_dir / "png") if a.render else None
    report = out.with_suffix(".report.md")
    write_report(report, b, plan, tokens, variant, out, fonts_note)

    print(f"{out}: {len(b.pages)} slides (variant {variant}) · report {report}")
    for f in fit:
        flags = [k for k in ("split", "scale", "shrunk", "reflow", "issues") if f.get(k)]
        if flags:
            print(f"  {f['out']:>2} S{f['src']} {f['layout']}: " + "; ".join(f"{k}={f[k]}" for k in flags))
    if fonts_note:
        print("\n".join(fonts_note))
    if missing_fonts:
        print("measure fonts substituted:", ", ".join(sorted(missing_fonts)))


if __name__ == "__main__":
    main()
