# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "PyYAML>=6"]
# ///
"""Measure design tokens from sample decks (deterministic; no judgment).

    uv run extract_design.py samples/*.pptx --out design/ [--primary FILE] [--exclude design/exclude.yaml]

Writes into --out:
  design.json   every measured shape, normalized to the primary deck's slide size (EMU)
  summary.md    compact digest for writing DESIGN.md: palette, type scale, chrome, layout clusters
  assets/       pictures that belong to the design (backgrounds, logos, repeated images)
  base.pptx     primary deck with slides removed, unused layouts/masters pruned, excluded images removed
  base.potx     same, as a PowerPoint template

All coordinates and font sizes from non-primary decks are scaled to the primary slide size, so a
21pt title on a 10" slide is reported as ~28pt (raw values are kept alongside).
"""
from __future__ import annotations

import argparse
import colorsys
import collections
import hashlib
import io
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

EMU_IN = 914400
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"a": A, "p": P}
CHROME_MIN_SHARE = 0.75  # an element on >= 75% of a deck's slides (and >= 2) is chrome
CLUSTER_MIN_SIM = 0.5    # Jaccard similarity to join a layout cluster
DHASH_MAX_DIST = 6


def q(tag):
    pfx, local = tag.split(":")
    return f"{{{NS[pfx]}}}{local}"


def inch(v):
    return round(v / EMU_IN, 2)


# ---------------------------------------------------------------- images

def dhash(blob: bytes, n: int = 8) -> str | None:
    try:
        im = Image.open(io.BytesIO(blob)).convert("L").resize((n + 1, n), Image.LANCZOS)
    except Exception:
        return None  # EMF/WMF/SVG
    px = list(im.tobytes())
    bits = 0
    for r in range(n):
        for c in range(n):
            bits = (bits << 1) | (px[r * (n + 1) + c] > px[r * (n + 1) + c + 1])
    return f"{bits:016x}"


def image_px(blob: bytes):
    try:
        return list(Image.open(io.BytesIO(blob)).size)
    except Exception:
        return None


class Excluder:
    def __init__(self, path: Path | None):
        self.rules = []
        if path and path.exists():
            self.rules = (yaml.safe_load(path.read_text()) or {}).get("images", [])
        self.hits = collections.defaultdict(list)

    def match(self, md5: str, dh: str | None, where: str) -> str | None:
        for i, rule in enumerate(self.rules):
            md5s = rule.get("md5") or []
            hit = any(md5.startswith(m) for m in md5s)
            if not hit and dh and rule.get("dhash"):
                hit = bin(int(dh, 16) ^ int(rule["dhash"], 16)).count("1") <= DHASH_MAX_DIST
            if hit:
                self.hits[i].append(where)
                return rule.get("note", f"rule {i}")
        return None


# ---------------------------------------------------------------- theme / color

def apply_mods(hexv: str, clr_el) -> tuple[str, float]:
    r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
    alpha = 1.0
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    for m in clr_el:
        tag = etree.QName(m).localname
        val = int(m.get("val", "0")) / 100000
        if tag == "lumMod":
            l *= val
        elif tag == "lumOff":
            l += val
        elif tag == "alpha":
            alpha = val
        elif tag == "shade":
            l *= val
        elif tag == "tint":
            l = l * val + (1 - val)
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(c * 255):02X}" for c in (r, g, b)), round(alpha, 3)


class Theme:
    def __init__(self, master):
        theme_part = master.part.part_related_by(RT.THEME)
        el = etree.fromstring(theme_part.blob)
        self.name = el.get("name")
        self.colors = {}
        scheme = el.find(".//a:clrScheme", NS)
        for c in scheme if scheme is not None else []:
            v = c[0]
            self.colors[etree.QName(c).localname] = (v.get("val") if etree.QName(v).localname == "srgbClr"
                                                     else v.get("lastClr", "000000")).upper()
        cmap = master._element.find(q("p:clrMap"))
        self.clrmap = dict(cmap.attrib) if cmap is not None else {}
        self.fonts = {}
        for kind in ("major", "minor"):
            f = el.find(f".//a:{kind}Font", NS)
            if f is None:
                continue
            latin, ea = f.find("a:latin", NS), f.find("a:ea", NS)
            hang = f.find("a:font[@script='Hang']", NS)
            self.fonts[kind] = {
                "latin": latin.get("typeface") if latin is not None else "",
                "ea": (ea.get("typeface") if ea is not None else "") or (hang.get("typeface") if hang is not None else ""),
            }

    def font(self, typeface: str | None) -> str | None:
        if not typeface or not typeface.startswith("+"):
            return typeface
        kind = "major" if typeface.startswith("+mj") else "minor"
        script = "ea" if typeface.endswith("-ea") else "latin"
        return self.fonts.get(kind, {}).get(script) or None

    def color(self, holder) -> dict | None:
        """holder: a:solidFill or any element whose first child is a color element."""
        if holder is None or len(holder) == 0:
            return None
        c = holder[0]
        tag = etree.QName(c).localname
        if tag == "srgbClr":
            base = c.get("val").upper()
        elif tag == "sysClr":
            base = c.get("lastClr", "000000").upper()
        elif tag == "schemeClr":
            key = self.clrmap.get(c.get("val"), c.get("val"))
            base = self.colors.get(key)
            if base is None:
                return {"scheme": c.get("val")}
        elif tag == "prstClr":
            base = {"black": "000000", "white": "FFFFFF"}.get(c.get("val"), "808080")
        else:
            return None
        hexv, alpha = apply_mods(base, c)
        out = {"hex": hexv}
        if alpha < 1:
            out["alpha"] = alpha
        if tag == "schemeClr":
            out["scheme"] = c.get("val")
        return out


# ---------------------------------------------------------------- text style resolution

def lvl_ppr(lst, level):
    if lst is None:
        return None
    return lst.find(f"a:lvl{level + 1}pPr", NS)


def text_chain(shape, theme_ctx, prs_default):
    """lstStyle-like elements to consult for inherited text properties, nearest first."""
    chain = []
    tb = shape._element.find(".//p:txBody", NS)
    if tb is not None:
        chain.append(tb.find("a:lstStyle", NS))
    if shape.is_placeholder:
        base = getattr(shape, "_base_placeholder", None)
        while base is not None:
            btb = base._element.find(".//p:txBody", NS)
            if btb is not None:
                chain.append(btb.find("a:lstStyle", NS))
            base = getattr(base, "_base_placeholder", None)
        ph_type = str(shape.placeholder_format.type)
        styles = theme_ctx["master_el"].find("p:txStyles", NS)
        if styles is not None:
            key = "p:titleStyle" if ("TITLE" in ph_type) else "p:bodyStyle" if any(
                k in ph_type for k in ("BODY", "OBJECT", "SUBTITLE")) else "p:otherStyle"
            chain.append(styles.find(key, NS))
    else:
        chain.append(prs_default)
        styles = theme_ctx["master_el"].find("p:txStyles", NS)
        if styles is not None:
            chain.append(styles.find("p:otherStyle", NS))
    return [c for c in chain if c is not None]


def first_attr(elems, attr):
    for e in elems:
        if e is not None and e.get(attr) is not None:
            return e.get(attr)
    return None


def first_child(elems, path):
    for e in elems:
        if e is not None:
            c = e.find(path, NS)
            if c is not None:
                return c
    return None


def read_text(shape, theme: Theme, theme_ctx, prs_default, scale):
    tf = shape.text_frame
    chain = text_chain(shape, theme_ctx, prs_default)
    body = tf._txBody.find("a:bodyPr", NS)
    paras = []
    for p in tf.paragraphs:
        ppr = p._p.find("a:pPr", NS)
        level = int(ppr.get("lvl", "0")) if ppr is not None else 0
        pprs = [ppr] + [lvl_ppr(c, level) for c in chain]
        defrprs = [x.find("a:defRPr", NS) for x in pprs if x is not None]
        bullet = None
        if first_child(pprs, "a:buNone") is not None and (ppr is not None and ppr.find("a:buNone", NS) is not None):
            bullet = None
        else:
            bc = first_child(pprs, "a:buChar")
            ba = first_child(pprs, "a:buAutoNum")
            bn = first_child(pprs, "a:buNone")
            # nearest definition wins
            for x in pprs:
                if x is None:
                    continue
                if x.find("a:buNone", NS) is not None:
                    bullet = None
                    break
                if x.find("a:buChar", NS) is not None:
                    bullet = x.find("a:buChar", NS).get("char")
                    break
                if x.find("a:buAutoNum", NS) is not None:
                    bullet = "autonum:" + x.find("a:buAutoNum", NS).get("type", "")
                    break
            del bc, ba, bn
        spc = {}
        for name in ("lnSpc", "spcBef", "spcAft"):
            el = first_child(pprs, f"a:{name}")
            if el is not None and len(el):
                v = el[0]
                spc[name] = (f"{int(v.get('val')) / 1000:g}%" if etree.QName(v).localname == "spcPct"
                             else f"{int(v.get('val')) / 100 * scale:g}pt")
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            rpr = r._r.find("a:rPr", NS)
            rprs = [rpr] + defrprs
            latin = first_child(rprs, "a:latin")
            ea = first_child(rprs, "a:ea")
            fill = first_child(rprs, "a:solidFill")
            sz = first_attr(rprs, "sz")
            runs.append({
                "text": r.text,
                "font": theme.font(latin.get("typeface") if latin is not None else "+mn-lt"),
                "ea": theme.font(ea.get("typeface") if ea is not None else "+mn-ea"),
                "size": round(int(sz) / 100 * scale, 1) if sz else None,
                "size_raw": int(sz) / 100 if sz else None,
                "bold": first_attr(rprs, "b") in ("1", "true"),
                "italic": first_attr(rprs, "i") in ("1", "true"),
                "color": theme.color(fill) if fill is not None else None,
                "caps": first_attr(rprs, "cap"),
                "spc": first_attr(rprs, "spc"),
            })
        paras.append({
            "level": level,
            "align": first_attr(pprs, "algn") or "l",
            "bullet": bullet,
            "marL": round(int(first_attr(pprs, "marL") or 0) * scale),
            "indent": round(int(first_attr(pprs, "indent") or 0) * scale),
            **spc,
            "runs": runs,
        })
    ins = {k: round(int(body.get(k)) * scale) for k in ("lIns", "tIns", "rIns", "bIns") if body is not None and body.get(k)}
    autofit = None
    if body is not None:
        for t in ("normAutofit", "spAutoFit", "noAutofit"):
            if body.find(f"a:{t}", NS) is not None:
                autofit = t
    return {
        "anchor": body.get("anchor", "t") if body is not None else "t",
        "wrap": body.get("wrap", "square") if body is not None else "square",
        "insets": ins,
        "autofit": autofit,
        "paragraphs": paras,
        "chars": sum(len(r["text"]) for pp in paras for r in pp["runs"]),
    }


# ---------------------------------------------------------------- shape walk

class Ctx:
    def __init__(self, deck_idx, prs, scale_x, scale_y, excluder, assets_dir, prs_default):
        self.deck_idx, self.prs = deck_idx, prs
        self.sx, self.sy = scale_x, scale_y
        self.excluder, self.assets_dir = excluder, assets_dir
        self.prs_default = prs_default
        self.blobs = {}  # md5 -> (blob, ext)


def fill_of(shape, theme):
    sppr = shape._element.find(q("p:spPr"))
    if sppr is None:
        return None
    for child in sppr:
        tag = etree.QName(child).localname
        if tag == "solidFill":
            return theme.color(child)
        if tag == "gradFill":
            stops = [theme.color(gs) for gs in child.findall(".//a:gs", NS)]
            return {"gradient": [s for s in stops if s]}
        if tag == "blipFill":
            return {"image": True}
        if tag == "noFill":
            return None
    style = shape._element.find(q("p:style"))
    if style is not None:
        fr = style.find("a:fillRef", NS)
        if fr is not None and fr.get("idx") not in (None, "0") and len(fr):
            return theme.color(fr)
    return None


def line_of(shape, theme):
    sppr = shape._element.find(q("p:spPr"))
    ln = sppr.find("a:ln", NS) if sppr is not None else None
    color = None
    if ln is not None:
        if ln.find("a:noFill", NS) is not None:
            return None
        sf = ln.find("a:solidFill", NS)
        color = theme.color(sf) if sf is not None else None
    if color is None:
        style = shape._element.find(q("p:style"))
        lr = style.find("a:lnRef", NS) if style is not None else None
        if lr is not None and lr.get("idx") not in (None, "0") and len(lr):
            color = theme.color(lr)
    if color is None:
        return None
    w = int(ln.get("w")) if ln is not None and ln.get("w") else 9525
    dash = ln.find("a:prstDash", NS) if ln is not None else None
    return {**color, "width_pt": round(w / 12700, 2), "dash": dash.get("val") if dash is not None else "solid"}


def walk(shapes, ctx: Ctx, theme: Theme, theme_ctx, where: str, xf=(1.0, 1.0, 0.0, 0.0), depth=0):
    sx, sy, tx, ty = xf
    for sh in shapes:
        if sh.left is None and not sh.is_placeholder:
            continue
        try:
            x, y, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if x is None:
            continue
        X, Y, W, H = x * sx + tx, y * sy + ty, w * sx, h * sy
        bbox = [round(X * ctx.sx), round(Y * ctx.sy), round(W * ctx.sx), round(H * ctx.sy)]
        rec = {"name": sh.name, "bbox": bbox, "depth": depth, "rot": round(sh.rotation, 1) if hasattr(sh, "rotation") else 0}
        if sh.is_placeholder:
            rec["ph"] = str(sh.placeholder_format.type).split(" ")[0].split(".")[-1]
            rec["ph_idx"] = sh.placeholder_format.idx
        st = sh.shape_type
        gd = sh._element.find(".//a:graphicData", NS)
        uri = gd.get("uri", "") if gd is not None else ""
        if st == MSO_SHAPE_TYPE.GROUP:
            rec["kind"] = "group"
            g = sh._element.find(q("p:grpSpPr")).find("a:xfrm", NS)
            off, ext = g.find("a:off", NS), g.find("a:ext", NS)
            choff, chext = g.find("a:chOff", NS), g.find("a:chExt", NS)
            csx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
            csy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
            ctx_x = int(off.get("x")) - int(choff.get("x")) * csx
            ctx_y = int(off.get("y")) - int(choff.get("y")) * csy
            child_xf = (sx * csx, sy * csy, tx + ctx_x * sx, ty + ctx_y * sy)
            rec["children"] = list(walk(sh.shapes, ctx, theme, theme_ctx, where, child_xf, depth + 1))
            yield rec
            continue
        if "diagram" in uri:
            rec["kind"] = "smartart"
        elif "chart" in uri:
            rec["kind"] = "chart"
        elif getattr(sh, "has_table", False) and sh.has_table:
            rec["kind"] = "table"
            tbl = sh.table
            rec["table"] = {"rows": len(tbl.rows), "cols": len(tbl.columns),
                            "col_w": [round(c.width * ctx.sx) for c in tbl.columns]}
            cells = []
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row.cells):
                    tcpr = cell._tc.find("a:tcPr", NS)
                    sf = tcpr.find("a:solidFill", NS) if tcpr is not None else None
                    first = None
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            if r.text.strip():
                                first = r
                                break
                        if first:
                            break
                    cells.append({"r": ri, "c": ci, "fill": theme.color(sf) if sf is not None else None,
                                  "text": cell.text[:40],
                                  "font": theme.font(first.font.name) if first else None,
                                  "size": round(first.font.size.pt * ctx.sx, 1) if first and first.font.size else None,
                                  "bold": bool(first and first.font.bold)})
            rec["table"]["cells"] = cells[:60]
        elif "ole" in uri.lower():
            rec["kind"] = "ole"
        elif st == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image") and st != MSO_SHAPE_TYPE.PLACEHOLDER:
            try:
                blob, ext = sh.image.blob, sh.image.ext
            except Exception:
                rec["kind"] = "picture-unreadable"
                yield rec
                continue
            md5 = hashlib.md5(blob).hexdigest()[:8]
            dh = dhash(blob)
            excl = ctx.excluder.match(md5, dh, f"{where}:{sh.name}")
            rec["kind"] = "picture"
            rec["picture"] = {"md5": md5, "dhash": dh, "ext": ext, "px": image_px(blob),
                              "crop": [round(v, 4) for v in (sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom)]}
            if excl:
                rec["excluded"] = excl
            else:
                ctx.blobs.setdefault(md5, (blob, ext))
        elif st in (MSO_SHAPE_TYPE.LINE,) or sh._element.tag == q("p:cxnSp"):
            rec["kind"] = "line"
            rec["line"] = line_of(sh, theme)
        elif st == MSO_SHAPE_TYPE.FREEFORM:
            rec["kind"] = "freeform"
            rec["fill"], rec["line"] = fill_of(sh, theme), line_of(sh, theme)
        else:
            geom = sh._element.find(".//a:prstGeom", NS)
            rec["geom"] = geom.get("prst") if geom is not None else None
            rec["fill"], rec["line"] = fill_of(sh, theme), line_of(sh, theme)
            has_text = sh.has_text_frame and sh.text_frame.text.strip()
            if has_text:
                rec["kind"] = "text"
                rec["text"] = read_text(sh, theme, theme_ctx, ctx.prs_default, ctx.sx)
            elif rec["fill"] or rec["line"]:
                rec["kind"] = "shape"
            elif sh.is_placeholder:
                rec["kind"] = "placeholder"
            else:
                rec["kind"] = "empty"
        if rec.get("kind") in ("text", "shape") and sh.has_text_frame and not rec.get("text") and sh.is_placeholder:
            rec["text"] = read_text(sh, theme, theme_ctx, ctx.prs_default, ctx.sx)
        yield rec


def flatten(recs):
    for r in recs:
        yield r
        yield from flatten(r.get("children", []))


# ---------------------------------------------------------------- per-deck extraction

def bg_of(el, theme):
    bg = el.find(".//p:cSld/p:bg", NS)
    if bg is None:
        return None
    bgpr = bg.find("p:bgPr", NS)
    if bgpr is not None:
        sf = bgpr.find("a:solidFill", NS)
        if sf is not None:
            return {"solid": theme.color(sf)}
        if bgpr.find("a:blipFill", NS) is not None:
            return {"image": True}
        if bgpr.find("a:gradFill", NS) is not None:
            return {"gradient": True}
    ref = bg.find("p:bgRef", NS)
    if ref is not None:
        return {"ref": ref.get("idx"), "color": theme.color(ref)}
    return None


def extract_deck(deck_idx, path: Path, canon, excluder, assets_dir):
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    sx, sy = canon[0] / W, canon[1] / H
    prs_default = prs.part._element.find("p:defaultTextStyle", NS)
    ctx = Ctx(deck_idx, prs, sx, sy, excluder, assets_dir, prs_default)
    deck = {"file": path.name, "size_emu": [W, H], "size_in": [inch(W), inch(H)], "scale": round(sx, 5),
            "masters": [], "slides": []}
    used_layouts = collections.Counter(id(s.slide_layout) for s in prs.slides)
    themes = {}
    for mi, m in enumerate(prs.slide_masters):
        th = Theme(m)
        themes[id(m)] = th
        tctx = {"master_el": m._element}
        mrec = {"index": mi, "theme": th.name, "theme_colors": th.colors, "theme_fonts": th.fonts,
                "bg": bg_of(m._element, th),
                "shapes": list(walk(m.shapes, ctx, th, tctx, f"{path.name}:M{mi}")),
                "layouts": []}
        for li, lay in enumerate(m.slide_layouts):
            mrec["layouts"].append({
                "index": li, "name": lay.name, "used_by": used_layouts.get(id(lay), 0),
                "bg": bg_of(lay._element, th),
                "shapes": list(walk(lay.shapes, ctx, th, tctx, f"{path.name}:M{mi}L{li}")),
            })
        deck["masters"].append(mrec)
    for si, s in enumerate(prs.slides, 1):
        m = s.slide_layout.slide_master
        th = themes[id(m)]
        tctx = {"master_el": m._element}
        mi = list(prs.slide_masters).index(m)
        li = list(m.slide_layouts).index(s.slide_layout)
        deck["slides"].append({
            "n": si, "layout": f"M{mi}L{li}:{s.slide_layout.name}",
            "bg": bg_of(s._element, th),
            "shapes": list(walk(s.shapes, ctx, th, tctx, f"{path.name}:S{si}")),
            "notes": bool(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()),
        })
    return deck, ctx.blobs


# ---------------------------------------------------------------- analysis

def rnd(v, step):
    return int(round(v / step))


def text_sig(rec):
    runs = [r for p in rec["text"]["paragraphs"] for r in p["runs"] if r["text"].strip()]
    if not runs:
        return None
    r = runs[0]
    return (r["font"], r["size"], r["bold"], (r["color"] or {}).get("hex"))


def chrome_key(rec):
    b = rec["bbox"]
    pos = tuple(rnd(v, EMU_IN * 0.05) for v in b)
    k = rec["kind"]
    if k == "picture":
        return ("picture", rec["picture"]["md5"], pos)
    if k == "text":
        return ("text", text_sig(rec), pos)
    if k in ("shape", "line"):
        return (k, rec.get("geom"), (rec.get("fill") or {}).get("hex"), (rec.get("line") or {}).get("hex"), pos)
    return None


def find_chrome(deck):
    n = len(deck["slides"])
    counts = collections.Counter()
    samples = collections.defaultdict(list)
    for s in deck["slides"]:
        seen = set()
        for rec in s["shapes"]:
            key = chrome_key(rec)
            if key and key not in seen:
                seen.add(key)
                counts[key] += 1
                samples[key].append((s["n"], rec))
    need = max(2, CHROME_MIN_SHARE * n)
    chrome = []
    for key, c in counts.items():
        if c < need:
            continue
        recs = [r for _, r in samples[key]]
        rec = recs[0]
        texts = sorted({"".join(rr["text"] for p in r["text"]["paragraphs"] for rr in p["runs"]) for r in recs}) if rec["kind"] == "text" else []
        chrome.append({"key": key, "slides": [n_ for n_, _ in samples[key]], "count": c, "kind": rec["kind"],
                       "rec": rec, "static_text": len(texts) == 1, "examples": texts[:4]})
    chrome.sort(key=lambda c: (c["rec"]["bbox"][1], c["rec"]["bbox"][0]))
    keys = {c["key"] for c in chrome}
    for s in deck["slides"]:
        s["chrome"] = [i for i, c in enumerate(chrome) if s["n"] in c["slides"]]
        for rec in s["shapes"]:
            rec["is_chrome"] = chrome_key(rec) in keys
    return chrome


def content_tokens(slide):
    """Column structure of the content: (kind, x, width) on a 0.5in grid. y and height are ignored
    because text boxes in the same layout grow/shrink and shift with their text."""
    toks = set()
    for rec in slide["shapes"]:
        if rec.get("is_chrome") or rec.get("excluded") or rec["kind"] in ("empty", "placeholder"):
            continue
        b = rec["bbox"]
        toks.add((rec["kind"], rnd(b[0], EMU_IN * 0.5), rnd(b[2], EMU_IN * 0.5)))
    return toks


def jaccard(a, b):
    union = len(a | b)
    return len(a & b) / union if union else 1.0


def cluster_slides(decks):
    clusters = []
    for d in decks:
        for s in d["slides"]:
            toks = content_tokens(s)
            best, best_sim = None, 0
            for c in clusters:
                if c["deck"] != d["file"]:
                    continue
                sim = jaccard(toks, c["tokens"])
                if sim > best_sim:
                    best, best_sim = c, sim
            if best and best_sim >= CLUSTER_MIN_SIM:
                best["members"].append(s["n"])
            else:
                clusters.append({"deck": d["file"], "rep": s["n"], "members": [s["n"]], "tokens": toks})
    for i, c in enumerate(clusters):
        c["id"] = f"C{i + 1}"
        del c["tokens"]
    return clusters


def walk_all_text(decks):
    for d in decks:
        for s in d["slides"]:
            for rec in flatten(s["shapes"]):
                if rec.get("kind") == "text":
                    for p in rec["text"]["paragraphs"]:
                        for r in p["runs"]:
                            if r["text"].strip():
                                yield d, s, rec, p, r


def stats(decks):
    type_scale = collections.Counter()
    text_colors = collections.Counter()
    fonts = collections.Counter()
    for d, s, rec, p, r in walk_all_text(decks):
        n = len(r["text"].strip())
        type_scale[(r["font"], r["size"], r["bold"], (r["color"] or {}).get("hex"))] += n
        if r["color"] and r["color"].get("hex"):
            text_colors[r["color"]["hex"]] += n
        fonts[r["font"]] += n
        if r["ea"] and any(ord(ch) > 0x3000 for ch in r["text"]):
            fonts[f"{r['ea']} (ea)"] += n
    fills = collections.Counter()
    lines = collections.Counter()
    canon_area = None
    for d in decks:
        for s in d["slides"]:
            for rec in flatten(s["shapes"]):
                f = rec.get("fill")
                if f and f.get("hex"):
                    area = rec["bbox"][2] * rec["bbox"][3]
                    fills[(f["hex"], f.get("alpha", 1))] += area
                ln = rec.get("line")
                if ln and ln.get("hex"):
                    lines[(ln["hex"], ln["width_pt"], ln.get("dash"))] += 1
    return type_scale, text_colors, fonts, fills, lines


def installed_fonts(extra_dirs):
    names = set()
    try:
        out = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True, timeout=60).stdout
        for line in out.splitlines():
            names.update(n.strip().lower() for n in line.split(","))
        for d in extra_dirs:
            for f in Path(d).glob("*.[ot]tf"):
                o = subprocess.run(["fc-scan", "--format", "%{family}\n", str(f)], capture_output=True, text=True).stdout
                names.update(n.strip().lower() for n in o.replace("\n", ",").split(",") if n.strip())
    except FileNotFoundError:
        return None
    return names


# ---------------------------------------------------------------- base template

def build_base(primary: Path, out_dir: Path, excluder: Excluder):
    prs = Presentation(str(primary))
    used = {id(s.slide_layout) for s in prs.slides}
    sld_ids = prs.slides._sldIdLst
    for sid in list(sld_ids):
        prs.part.drop_rel(sid.get(f"{{{NS_R}}}id"))
        sld_ids.remove(sid)
    kept = []
    for m in list(prs.slide_masters):
        for lay in list(m.slide_layouts):
            if id(lay) not in used:
                m.slide_layouts.remove(lay)
        if len(m.slide_layouts) == 0:
            lst = prs.part._element.find("p:sldMasterIdLst", NS)
            for mid in list(lst):
                if prs.part.related_part(mid.get(f"{{{NS_R}}}id")) is m.part:
                    prs.part.drop_rel(mid.get(f"{{{NS_R}}}id"))
                    lst.remove(mid)
        else:
            kept.append(m)
    removed = []
    for m in kept:
        for owner in [m, *m.slide_layouts]:
            for sh in list(owner.shapes):
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    blob = sh.image.blob
                    if excluder.match(hashlib.md5(blob).hexdigest()[:8], dhash(blob), f"base:{owner.name}"):
                        sh._element.getparent().remove(sh._element)
                        removed.append(f"{getattr(owner, 'name', 'master')}:{sh.name}")
    base = out_dir / "base.pptx"
    prs.save(str(base))
    potx = out_dir / "base.potx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(potx, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            zout.writestr(item, data)
    layouts = [(mi, m.name if hasattr(m, "name") else "", [l.name for l in m.slide_layouts]) for mi, m in enumerate(kept)]
    return layouts, removed


NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ---------------------------------------------------------------- summary

def fmt_color(c):
    if not c:
        return "-"
    s = c.get("hex", c.get("scheme", "?"))
    if c.get("alpha") is not None and c.get("alpha") < 1:
        s += f"@{c['alpha']}"
    return s


def describe(rec):
    b = rec["bbox"]
    pos = f"({inch(b[0])}, {inch(b[1])}) {inch(b[2])}×{inch(b[3])}in"
    k = rec["kind"]
    if k == "text":
        sig = text_sig(rec) or (None, None, False, None)
        txt = " / ".join("".join(r["text"] for r in p["runs"]) for p in rec["text"]["paragraphs"])[:60]
        bul = {p["bullet"] for p in rec["text"]["paragraphs"]} - {None}
        al = {p["align"] for p in rec["text"]["paragraphs"]}
        extra = f" fill={fmt_color(rec.get('fill'))}" if rec.get("fill") else ""
        return (f"text {pos} {sig[0]} {sig[1]}pt{' B' if sig[2] else ''} #{sig[3]} align={','.join(al)}"
                f"{' bullets=' + ','.join(bul) if bul else ''}{extra} paras={len(rec['text']['paragraphs'])} chars={rec['text']['chars']} «{txt}»")
    if k == "picture":
        pic = rec["picture"]
        return f"picture {pos} {pic['md5']} {pic['px']}{' EXCLUDED' if rec.get('excluded') else ''}"
    if k in ("shape", "line", "freeform"):
        return f"{k} {pos} geom={rec.get('geom')} fill={fmt_color(rec.get('fill'))} line={fmt_color(rec.get('line'))}{' ' + str(rec['line']['width_pt']) + 'pt' if rec.get('line') else ''}"
    if k == "table":
        return f"table {pos} {rec['table']['rows']}×{rec['table']['cols']}"
    if k == "group":
        kinds = collections.Counter(c["kind"] for c in flatten(rec.get("children", [])))
        return f"group {pos} {dict(kinds)}"
    return f"{k} {pos}"


def sample_ids(design_md=Path("DESIGN.md")):
    """file name -> sample id (C, R …) from DESIGN.md meta.samples, so evidence matches the tokens."""
    try:
        from design_tokens import load
        return {smp["file"]: smp["id"] for smp in load(design_md)["meta"]["samples"]}
    except Exception:
        return {}


def box_styles(decks, canon):
    """Box (shape) styles used on slides: (geom, fill, line, dash) with evidence, typical size and the text
    style drawn inside — the shape's own text, or a text box lying on the same rectangle (Google Slides
    decks draw labels as separate boxes). Full-slide backgrounds and lines are skipped."""
    tol = 0.06 * 914400
    out = {}
    ids = sample_ids()
    for d in decks:
        did = ids.get(d["file"], d["file"][0].upper())
        for s in d["slides"]:
            recs = list(flatten(s["shapes"]))
            texts = [r for r in recs if r.get("kind") == "text"]
            for r in recs:
                if r.get("kind") not in ("shape", "text") or not r.get("geom"):
                    continue
                f, ln = r.get("fill"), r.get("line")
                if not (f and f.get("hex")) and not (ln and ln.get("hex")):
                    continue
                x, y, w, h = r["bbox"]
                if w * h > 0.8 * canon[0] * canon[1]:
                    continue
                key = (r["geom"], (f or {}).get("hex"), (f or {}).get("alpha", 1) if f else None,
                       (ln or {}).get("hex"), (ln or {}).get("width_pt"), (ln or {}).get("dash"))
                e = out.setdefault(key, {"count": 0, "evidence": [], "sizes": [], "text": collections.Counter()})
                e["count"] += 1
                tag = f"{did}:S{s['n']}"
                if tag not in e["evidence"]:
                    e["evidence"].append(tag)
                e["sizes"].append((w, h))
                inner = [r] if r.get("kind") == "text" else [
                    t for t in texts if t is not r and t["bbox"][0] >= x - tol and t["bbox"][1] >= y - tol
                    and t["bbox"][0] + t["bbox"][2] <= x + w + tol and t["bbox"][1] + t["bbox"][3] <= y + h + tol]
                for t in inner:
                    sig = text_sig(t)
                    if sig:
                        e["text"][sig] += 1
    res = []
    for (geom, fhex, falpha, lhex, lw, dash), e in sorted(out.items(), key=lambda kv: -kv[1]["count"]):
        ws, hs = sorted(w for w, _ in e["sizes"]), sorted(h for _, h in e["sizes"])
        res.append({"geom": geom, "fill": fhex, "fill_alpha": falpha, "line": lhex, "line_pt": lw, "dash": dash,
                    "count": e["count"], "evidence": e["evidence"],
                    "median_in": [inch(ws[len(ws) // 2]), inch(hs[len(hs) // 2])],
                    "text": [{"font": k[0], "size": k[1], "bold": k[2], "color": k[3], "n": n} for k, n in e["text"].most_common(3)]})
    return res


def write_summary(out: Path, decks, chrome_by_deck, clusters, st, font_avail, excluder, base_info, canon, boxes=()):
    type_scale, text_colors, fonts, fills, lines = st
    L = []
    L.append("# Design extraction summary\n")
    L.append(f"Canonical slide size: {inch(canon[0])} × {inch(canon[1])} in ({canon[0]} × {canon[1]} EMU). "
             "All values below are normalized to it.\n")
    for d in decks:
        L.append(f"- `{d['file']}`: {d['size_in'][0]}×{d['size_in'][1]}in, {len(d['slides'])} slides, scale ×{d['scale']}")
    L.append("\n## Excluded elements\n")
    for i, rule in enumerate(excluder.rules):
        hits = excluder.hits.get(i, [])
        L.append(f"- {rule.get('note', '').strip()} — matched {len(hits)} placements")
    L.append("\n## Fonts (char-weighted)\n")
    total = sum(fonts.values()) or 1
    for f, n in fonts.most_common(15):
        name = (f or "?").replace(" (ea)", "")
        avail = "?" if font_avail is None else ("installed" if name.lower() in font_avail else "MISSING")
        L.append(f"- {f}: {n} chars ({n * 100 // total}%) — {avail}")
    L.append("\n## Type scale (font, size pt, bold, color) — char-weighted, top 25\n")
    for (f, s, b, c), n in type_scale.most_common(25):
        L.append(f"- {f} {s}pt{' bold' if b else ''} #{c}: {n}")
    L.append("\n## Text colors\n")
    for c, n in text_colors.most_common(12):
        L.append(f"- #{c}: {n} chars")
    L.append("\n## Fill colors (area-weighted, in slide areas)\n")
    slide_area = canon[0] * canon[1]
    for (c, a), n in fills.most_common(12):
        L.append(f"- #{c}{'@' + str(a) if a != 1 else ''}: {n / slide_area:.2f}")
    L.append("\n## Line colors\n")
    for (c, w, dash), n in lines.most_common(10):
        L.append(f"- #{c} {w}pt {dash}: {n}")
    L.append("\n## Box styles (geom, fill, line) — slides only, with text drawn inside\n")
    for bx in boxes:
        fill = f"#{bx['fill']}{'@' + str(bx['fill_alpha']) if bx['fill_alpha'] not in (None, 1) else ''}" if bx["fill"] else "-"
        line = f"#{bx['line']} {bx['line_pt']}pt{' ' + bx['dash'] if bx['dash'] and bx['dash'] != 'solid' else ''}" if bx["line"] else "-"
        txt = "; ".join(f"{t['font']} {t['size']}pt{' B' if t['bold'] else ''} #{t['color']}" for t in bx["text"]) or "no text"
        L.append(f"- {bx['geom']} fill={fill} line={line} ×{bx['count']} ~{bx['median_in'][0]}×{bx['median_in'][1]}in "
                 f"[{', '.join(bx['evidence'][:6])}] text: {txt}")
    L.append("\n## Masters / layouts actually used\n")
    for d in decks:
        for m in d["masters"]:
            for lay in m["layouts"]:
                if lay["used_by"]:
                    L.append(f"- {d['file']} M{m['index']}L{lay['index']} '{lay['name']}' used_by={lay['used_by']} "
                             f"theme='{m['theme']}' master_bg={m['bg']} layout_bg={lay['bg']}")
                    for rec in m["shapes"] + lay["shapes"]:
                        if rec["kind"] not in ("placeholder", "empty"):
                            L.append(f"    - {describe(rec)}")
    L.append("\n## Chrome (elements repeated on most slides of a deck)\n")
    for d in decks:
        L.append(f"### {d['file']}\n")
        for i, c in enumerate(chrome_by_deck[d["file"]]):
            tag = "static" if c["static_text"] or c["kind"] != "text" else "SLOT"
            L.append(f"- [{i}] {tag} on {c['count']}/{len(d['slides'])}: {describe(c['rec'])}")
            if c["kind"] == "text" and not c["static_text"]:
                L.append(f"    examples: {c['examples'][:3]}")
    L.append("\n## Layout clusters (content shapes only; chrome removed)\n")
    for c in clusters:
        d = next(x for x in decks if x["file"] == c["deck"])
        s = d["slides"][c["rep"] - 1]
        missing = [i for i in range(len(chrome_by_deck[d["file"]])) if i not in s["chrome"]]
        L.append(f"### {c['id']} — {c['deck']} slides {c['members']} (rep S{c['rep']}, layout {s['layout']})")
        if missing:
            L.append(f"chrome missing on rep: {missing}")
        for rec in s["shapes"]:
            if rec.get("is_chrome") or rec["kind"] in ("empty",):
                continue
            L.append(f"- {describe(rec)}{'  [EXCLUDED]' if rec.get('excluded') else ''}")
        L.append("")
    L.append("## Base template\n")
    L.append(f"- kept masters/layouts: {base_info[0]}")
    L.append(f"- excluded pictures removed from masters/layouts: {base_info[1] or 'none'}")
    out.write_text("\n".join(L) + "\n")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("samples", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, default=Path("design"))
    ap.add_argument("--primary", type=Path, help="deck whose slide size/masters are canonical (default: most slides)")
    ap.add_argument("--exclude", type=Path, help="exclude.yaml (default: <out>/exclude.yaml)")
    ap.add_argument("--font-dir", action="append", default=[], help="extra font dir to count as installed (default ./.fonts)")
    a = ap.parse_args()

    samples = [p for p in a.samples if p.suffix.lower() == ".pptx"]
    skipped = [p for p in a.samples if p not in samples]
    if skipped:
        print(f"skipping non-pptx: {[str(s) for s in skipped]}", file=sys.stderr)
    if not samples:
        sys.exit("no .pptx samples")
    a.out.mkdir(parents=True, exist_ok=True)
    assets = a.out / "assets"
    if assets.exists():
        shutil.rmtree(assets)
    assets.mkdir()
    excluder = Excluder(a.exclude or a.out / "exclude.yaml")

    primary = a.primary or max(samples, key=lambda p: len(Presentation(str(p)).slides))
    pp = Presentation(str(primary))
    canon = (pp.slide_width, pp.slide_height)

    decks, blobs = [], {}
    for i, path in enumerate(samples):
        d, b = extract_deck(i, path, canon, excluder, assets)
        ar = d["size_emu"][0] / d["size_emu"][1]
        if abs(ar - canon[0] / canon[1]) > 0.01:
            print(f"warning: {path.name} aspect {ar:.3f} differs from primary; coordinates scaled per axis", file=sys.stderr)
        decks.append(d)
        blobs.update(b)

    chrome_by_deck = {d["file"]: find_chrome(d) for d in decks}
    clusters = cluster_slides(decks)
    st = stats(decks)
    boxes = box_styles(decks, canon)

    # design assets: master/layout pictures, chrome pictures, pictures repeated on >= 2 slides
    keep = set()
    pic_slides = collections.Counter()
    for d in decks:
        for m in d["masters"]:
            for rec in flatten(m["shapes"] + [r for lay in m["layouts"] if lay["used_by"] for r in lay["shapes"]]):
                if rec.get("kind") == "picture" and not rec.get("excluded"):
                    keep.add(rec["picture"]["md5"])
        for s in d["slides"]:
            for md5 in {r["picture"]["md5"] for r in flatten(s["shapes"]) if r.get("kind") == "picture" and not r.get("excluded")}:
                pic_slides[md5] += 1
    keep |= {m for m, n in pic_slides.items() if n >= 2}
    for md5 in keep:
        if md5 in blobs:
            blob, ext = blobs[md5]
            (assets / f"{md5}.{ext}").write_bytes(blob)

    base_info = build_base(primary, a.out, excluder)
    font_dirs = a.font_dir or ([".fonts"] if Path(".fonts").is_dir() else [])
    font_avail = installed_fonts(font_dirs)

    data = {
        "canonical_size_emu": list(canon), "primary": primary.name,
        "decks": decks,
        "chrome": {k: [{**{kk: vv for kk, vv in c.items() if kk != "key"}} for c in v] for k, v in chrome_by_deck.items()},
        "clusters": clusters,
        "box_styles": boxes,
        "excluded": [{"note": r.get("note"), "placements": excluder.hits.get(i, [])} for i, r in enumerate(excluder.rules)],
        "assets": sorted(p.name for p in assets.iterdir()),
    }
    (a.out / "design.json").write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str))
    write_summary(a.out / "summary.md", decks, chrome_by_deck, clusters, st, font_avail, excluder, base_info, canon, boxes)
    print(f"wrote {a.out}/design.json, summary.md, base.pptx, base.potx, assets/ ({len(keep)} images)")
    print(f"clusters: {len(clusters)}; excluded placements: {sum(len(v) for v in excluder.hits.values())}")


if __name__ == "__main__":
    main()
